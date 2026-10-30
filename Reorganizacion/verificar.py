from pathlib import Path
import json, re, subprocess, sys, zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from pptx import Presentation
from docx import Document
import fitz

WORK=Path('/workspace/reorganizacion-unidades')
UNITS=json.loads((WORK/'catalogo.json').read_text())
VIEWS=WORK/'Vistas';VIEWS.mkdir(exist_ok=True)

def run(args):
    r=subprocess.run([str(a) for a in args],capture_output=True,text=True)
    if r.returncode:raise RuntimeError(r.stderr[-4500:])
    return r

def offices():
    unitarg=next((a.split('=',1)[1] for a in sys.argv if a.startswith('--unit=')),None)
    selected=[u for u in UNITS if not unitarg or u['id']==unitarg]
    files=[]
    for u in selected:
        files += [Path(u['folder'])/u['files'][kind] for kind in ('pptx','workshop')]
    result=run(['soffice','-env:UserInstallation=file:///tmp/lo-entrega-unidades','--headless','--convert-to','pdf','--outdir',VIEWS,*files])
    (WORK/'libreoffice.log').write_text(result.stdout+result.stderr)
    prior=WORK/'Verificacion_Office.json'
    report=[r for r in json.loads(prior.read_text()) if r['id']!=unitarg] if unitarg and prior.exists() else []
    for u in selected:
        prs=Presentation(Path(u['folder'])/u['files']['pptx'])
        assert len(prs.slides)==u['slide_count']
        offset=2
        for t in u['topics']:
            source=Presentation(Path(t['source'])/'Presentacion.pptx')
            for i,slide in enumerate(source.slides):
                actual=prs.slides[offset+i]
                assert len(actual.shapes)==len(slide.shapes),(u['id'],t['id'],i)
                assert [s.text for s in actual.shapes if s.has_text_frame]==[s.text for s in slide.shapes if s.has_text_frame],(u['id'],t['id'],i)
                assert t['title'] in actual.notes_slide.notes_text_frame.text
            offset+=len(source.slides)
        pdf=fitz.open(VIEWS/(Path(u['files']['pptx']).stem+'.pdf'))
        assert len(pdf)==len(prs.slides)
        assert all(len(page.get_text().strip())>40 for page in pdf)
        doc=Document(Path(u['folder'])/u['files']['workshop'])
        text='\n'.join(p.text for p in doc.paragraphs)
        for word in ('Taller grupal','Organización del equipo','revisión cruzada','Contribuciones','Rúbrica grupal','20 puntos'):
            assert word.lower() in text.lower(),(u['id'],word)
        assert 'Trabajo individual' not in text
        assert len(doc.tables)==(4 if u['id']=='PND_U4' else 3) and len(doc.tables[0].rows)==6 and len(doc.tables[-1].rows)==6
        assert sum(int(row.cells[-1].text) for row in doc.tables[-2].rows[1:])==20
        if u['id']=='PND_U4':
            assert 'Datos para resolver el taller' in text
            assert (Path(u['folder'])/'Recursos/PND_U4_4.4_Estudio_de_caso_Datos_simulados.csv').exists()
            assert len(doc.tables[1].rows)==5
            grouped={}
            for row in doc.tables[1].rows[1:]:
                period=row.cells[0].text
                n,h,correction,compliant=[int(c.text) for c in row.cells[1:]]
                totals=grouped.setdefault(period,[0,0,0,0])
                totals[0]+=n;totals[1]+=n*h;totals[2]+=correction;totals[3]+=compliant
            assert grouped=={'Inicial':[100,4000,12,80],'Futuro simulado':[100,2400,5,95]},grouped
        count=0
        for t in u['topics']:
            assert f"Sección {t['num']} · {t['title']}" in text
            for i in range(1,len(t['content']['tasks'])+1):
                assert f"Ejercicio {t['num']}.{i}" in text
                count+=1
        taskpdf=fitz.open(VIEWS/(Path(u['files']['workshop']).stem+'.pdf'))
        assert len(taskpdf)>=3
        tasktext=' '.join(page.get_text() for page in taskpdf)
        assert all(t['num'] in tasktext and t['title'].replace(' ','') in ''.join(tasktext.split()) for t in u['topics'])
        report.append(dict(id=u['id'],slides=len(prs.slides),all_original_slide_shapes_and_text_preserved=True,
                           workshop_pages=len(taskpdf),workshop_exercises=count,group_roles=5,
                           peer_review=True,contribution_record=True,rubric_points=20))
        print('Office comprobado:',u['id'],len(prs.slides),'diapositivas;',len(taskpdf),'páginas de taller;',count,'ejercicios',flush=True)
    (WORK/'Verificacion_Office.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))

def media_unit(u):
    checks={}
    for kind in ('video','audio'):
        file=Path(u['folder'])/u['files'][kind]
        info=json.loads(run(['ffprobe','-v','error','-show_streams','-show_format','-show_chapters','-of','json',file]).stdout)
        assert len(info['chapters'])==len(u['topics']),(u['id'],kind,'chapters',len(info['chapters']))
        for chapter,t in zip(info['chapters'],u['topics']):
            assert chapter['tags']['title'].startswith(t['num']) and t['title'] in chapter['tags']['title']
        if kind=='video':
            v=next(s for s in info['streams'] if s['codec_type']=='video')
            assert v['codec_name']=='h264' and v['width']==1920 and v['height']==1080
        decode=run(['ffmpeg','-v','error','-xerror','-threads','2','-i',file,'-map','0:v?','-map','0:a:0','-f','null','-'])
        assert not decode.stderr.strip(),(u['id'],kind,decode.stderr[-500:])
        checks[kind]=dict(seconds=float(info['format']['duration']),bytes=file.stat().st_size,
                          chapters=info['chapters'],fully_decoded_without_errors=True)
    target=WORK/u['id']/'Verificacion_medios.json'
    target.write_text(json.dumps(checks,ensure_ascii=False,indent=2))
    return u['id']

def browsers():
    from playwright.sync_api import sync_playwright
    from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
    from threading import Thread
    from urllib.parse import quote
    class Handler(SimpleHTTPRequestHandler):
        def __init__(self,*args,**kwargs):super().__init__(*args,directory='/workspace/Entrega',**kwargs)
        def log_message(self,*args):pass
        def handle(self):
            try:super().handle()
            except (BrokenPipeError,ConnectionResetError):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    Thread(target=server.serve_forever,daemon=True).start()
    def url(path):return 'http://127.0.0.1:'+str(server.server_port)+'/'+quote(path.relative_to('/workspace/Entrega').as_posix())
    results=[]
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1280,'height':900})
        for u in UNITS:
            errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto(url(Path(u['folder'])/u['files']['educaplay']))
            questions=page.evaluate('questions')
            assert len(questions)==u['questions']
            expected=[q for t in u['topics'] for q in t['questions']]
            assert questions==expected
            hrefs=page.locator('tbody a').evaluate_all('(els)=>els.map(e=>e.href)')
            assert hrefs==[t.get('activity_url',t['profile']) for t in u['topics']]
            for correct in (True,False):
                if not correct:page.locator('#restart').click()
                for q in questions:
                    answer=q['answer'] if correct else (q['answer']+1)%4
                    page.locator('#options button').nth(answer).click()
                    assert page.locator('#feedback').inner_text().endswith(q['feedback'])
                    assert page.locator('#options button:disabled').count()==4
                    page.locator('#next').click()
                assert page.locator('#result').inner_text()==f"Resultado: {len(questions) if correct else 0}/{len(questions)}"
            page.locator('#restart').click()
            page.locator('#options button').first.focus();page.keyboard.press('Enter')
            assert page.locator('#feedback').is_visible()
            page.set_viewport_size({'width':390,'height':844})
            overflow=page.evaluate('document.documentElement.scrollWidth>innerWidth')
            assert not overflow,u['id']
            page.set_viewport_size({'width':1280,'height':900})
            playback={}
            for kind,tag in [('video','video'),('audio','audio')]:
                uri=url(Path(u['folder'])/u['files'][kind])
                result=page.evaluate('''async ({uri,tag})=>{const media=document.createElement(tag);media.src=uri;media.muted=true;document.body.append(media);await new Promise((resolve,reject)=>{media.onloadedmetadata=resolve;media.onerror=()=>reject(new Error('media error'));});await media.play();await new Promise(r=>setTimeout(r,250));const out={duration:media.duration,currentTime:media.currentTime,error:media.error,paused:media.paused,width:media.videoWidth||null,height:media.videoHeight||null};media.pause();media.remove();return out}''',{'uri':uri,'tag':tag})
                assert result['currentTime']>0 and result['error'] is None and not result['paused']
                if kind=='video':assert result['width']==1920 and result['height']==1080
                playback[kind]=result
            assert not errors,errors
            results.append(dict(id=u['id'],questions=len(questions),all_correct=True,all_wrong=True,
                                keyboard=True,restart=True,mobile_no_horizontal_overflow=True,
                                activity_or_profile_links_match_received_assignment=True,
                                direct_activity_links_match_uploaded_pdf_and_syllabus=all(t.get('activity_url') for t in u['topics']),
                                remote_activity_correspondence_verified=False,
                                javascript_errors=errors,playback=playback))
            print('Navegador comprobado:',u['id'],len(questions),'preguntas y ambos medios reproducidos',flush=True)
        browser.close()
    server.shutdown();server.server_close()
    (WORK/'Verificacion_navegador.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))

if __name__=='__main__':
    if '--office' in sys.argv:offices()
    elif '--browser' in sys.argv:browsers()
    elif '--media' in sys.argv:
        with ProcessPoolExecutor(max_workers=2) as pool:
            futures={pool.submit(media_unit,u):u for u in UNITS}
            for f in as_completed(futures):print('Decodificación completa:',f.result(),flush=True)
