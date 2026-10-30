from pathlib import Path
import csv, hashlib, html, json, re, subprocess, sys, zipfile
from urllib.parse import unquote
from docx import Document
from docx.shared import Inches,Pt
import fitz

WORK=Path('/workspace/reorganizacion-unidades')
OUT=Path('/workspace/Entrega')
UNITS=json.loads((WORK/'catalogo.json').read_text())
ZIP=Path('/workspace/Entrega_Corregida_Por_Unidades.zip')

def stamp(seconds):
    n=round(float(seconds));return f'{n//60:02d}:{n%60:02d}'

def prepare():
    doc=Document()
    for s in doc.sections:
        s.top_margin=s.bottom_margin=Inches(.65)
        s.left_margin=s.right_margin=Inches(.7)
        s.footer.paragraphs[0].text='Proyecto educativo · Organización por materia, unidad y sección'
    doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(9.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(2)
    doc.styles['Normal'].paragraph_format.line_spacing=1
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
    doc.add_heading('Índice de la entrega por unidades',0)
    doc.add_paragraph('2 materias · 8 unidades · 27 secciones · cinco archivos principales por unidad')
    doc.add_paragraph('Cada unidad contiene una presentación PowerPoint, un video MP4, un audio MP3, un taller grupal Word y un archivo HTML de Educaplay con acceso por sección y práctica complementaria. Los nombres de los archivos indican materia, unidad, nombre de la unidad y rango de secciones. Abre 00_Indice.html para acceder directamente a cada archivo.')
    doc.add_paragraph('En Educaplay se recibieron seis enlaces a perfiles. El archivo de cada unidad indica qué perfil fue asignado a cada sección y lo distingue de un enlace directo a un juego. No se ha podido comprobar la correspondencia de los juegos remotos desde este entorno.')
    cards=[];previous=None;all_mapping=[]
    for u in UNITS:
        media=json.loads((WORK/u['id']/'media.json').read_text())
        if u['materia']!=previous:
            if previous is not None:doc.add_page_break()
            doc.add_heading(u['materia'],1);previous=u['materia']
        doc.add_heading(f"Unidad {u['number']}. {u['title']}",2)
        doc.add_paragraph(f"Presentación: {u['slide_count']} diapositivas. Video: {stamp(media['video']['format']['duration'])}. Audio: {stamp(media['audio']['format']['duration'])}.")
        rel=Path(u['folder']).relative_to(OUT)
        links=[]
        for key,label in [('pptx','Presentación PowerPoint'),('video','Video'),('audio','Audio'),('workshop','Taller grupal'),('educaplay','Educaplay y práctica')]:
            href=(rel/u['files'][key]).as_posix()
            assert (OUT/href).exists()
            links.append(f'<a href="{html.escape(href)}">{label}</a>')
        topics=[]
        for i,t in enumerate(u['topics']):
            slides=u['slides'][i];vc=media['chapters'][i];ac=media['audio_chapters'][i]
            detail=f"Diapositivas {slides['first']}–{slides['last']} · Video {stamp(vc['start'])} · Audio {stamp(ac['start'])}"
            doc.add_paragraph(f"{t['num']} · {t['title']}\n{detail}",style='List Bullet')
            topics.append('<li><strong>'+html.escape(t['num']+' · '+t['title'])+'</strong><br><small>'+html.escape(detail)+'</small></li>')
            all_mapping.append(dict(materia=u['materia'],unidad=u['number'],nombre_unidad=u['title'],seccion=t['num'],nombre_seccion=t['title'],carpeta=rel.as_posix(),archivos=u['files'],diapositivas=slides,video_inicio=vc['start'],audio_inicio=ac['start'],perfil_recibido=t['profile'],enlace_directo_verificado=False))
        cards.append(f'<article data-materia="{u["code"]}"><small>{html.escape(u["materia"])} · Unidad {u["number"]}</small><h2>{html.escape(u["title"])}</h2><p>{u["slide_count"]} diapositivas · Video {stamp(media["video"]["format"]["duration"])} · Audio {stamp(media["audio"]["format"]["duration"])}</p><nav>'+''.join(links)+'</nav><ol>'+''.join(topics)+'</ol></article>')
    doc.save(WORK/'00_Indice.docx')
    (OUT/'00_Indice.html').write_text('''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Entrega por unidades · dos materias</title><style>*{box-sizing:border-box}body{margin:0;font:17px/1.6 system-ui;color:#18394a;background:#f5f2eb}header{background:#112b3c;color:white;padding:45px max(24px,calc((100vw - 1200px)/2))}header p{max-width:1000px}h1{font-size:clamp(2rem,4vw,3rem);line-height:1.15}.wrap{max-width:1250px;margin:auto;padding:30px 24px}.filters{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:24px}button{font:inherit;border:2px solid #087e8b;border-radius:8px;background:white;padding:10px 15px;color:#18394a;cursor:pointer}button[aria-pressed=true]{background:#087e8b;color:white}.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:24px}article{padding:25px;background:white;border-radius:15px;border-top:5px solid #087e8b}article[hidden]{display:none}h2{line-height:1.25;font-size:1.4rem;overflow-wrap:anywhere}small{color:#536975}nav{display:flex;gap:8px;flex-wrap:wrap}a{color:#086977;overflow-wrap:anywhere}nav a{padding:9px 12px;background:#e2f0ee;border-radius:7px;font-weight:bold;text-decoration:none}li{margin:12px 0}li small{font-size:.85rem}a:focus-visible,button:focus-visible{outline:4px solid #b9432e;outline-offset:3px}.note{padding:20px;border-left:5px solid #e7ae53;background:#fff8e9;margin-top:28px}@media(max-width:760px){.grid{grid-template-columns:1fr}.wrap{padding:22px 15px}article{padding:20px}}</style></head><body><header><h1>Proyecto educativo<br>Entrega organizada por unidades</h1><p>2 materias · 8 unidades · 27 secciones. En cada unidad: una presentación, un video, un audio, un taller grupal y un archivo de acceso a Educaplay con práctica complementaria.</p></header><main class="wrap"><div class="filters"><button data-filter="all" aria-pressed="true">Todas las unidades</button><button data-filter="NE" aria-pressed="false">Negocios Electrónicos</button><button data-filter="PND" aria-pressed="false">Procesos para Negocios Digitales</button></div><section class="grid">'''+''.join(cards)+'''</section><aside class="note">Educaplay: los enlaces recibidos corresponden a seis perfiles. La correspondencia con los juegos individuales está pendiente de comprobación. Las preguntas de las 27 secciones están disponibles en las prácticas complementarias de cada unidad; estas funcionan localmente y son independientes de Educaplay.</aside></main><script>const cards=[...document.querySelectorAll('article')],buttons=[...document.querySelectorAll('button[data-filter]')];buttons.forEach(b=>b.onclick=()=>{cards.forEach(c=>c.hidden=b.dataset.filter!=='all'&&c.dataset.materia!==b.dataset.filter);buttons.forEach(x=>x.setAttribute('aria-pressed',String(x===b)));});</script></body></html>''',encoding='utf-8')
    (OUT/'LEEME.txt').write_text('''ENTREGA ORGANIZADA POR UNIDADES

1. Extrae todo el ZIP.
2. Abre 00_Indice.html para elegir una materia y una unidad.
3. También puedes abrir las carpetas: materia / unidad con su nombre completo.

CADA UNIDAD CONTIENE CINCO ARCHIVOS PRINCIPALES:
- Presentacion.pptx: material completo y editable de todas sus secciones.
- Video.mp4: explicación con capítulos identificados por número y nombre de sección.
- Audio.mp3: explicación de todas sus secciones, con capítulos identificados.
- Taller_grupal.docx: roles, casos y ejercicios por sección, revisión entre compañeros,
  producto común, defensa, contribuciones y rúbrica grupal de 20 puntos propuesta.
- Educaplay.html: acceso a los perfiles asignados por sección y práctica complementaria
  funcional con todas las preguntas preparadas de esa unidad.

Los nombres llevan materia abreviada, número y nombre de unidad y rango de secciones.
El índice indica dónde empieza cada sección en la presentación, el video y el audio.
La unidad 4 de Procesos incluye además Recursos con el CSV de datos simulados
del estudio de caso; el propio taller contiene un cuadro de datos suficiente para resolverlo.

EDUCAPLAY: se recibieron seis enlaces de perfiles, no los enlaces de cada juego.
La asociación de cada sección a una cuenta sigue el reparto comunicado por el estudiante.
La correspondencia y el contenido remoto de los juegos no se han podido comprobar desde
este entorno. Las prácticas HTML son locales e independientes de Educaplay.

Las voces son sintéticas neuronales en español. Los casos y datos son ficticios.
Los 27 temas se mantienen completos dentro de las ocho unidades del sílabo.
Los tiempos, equipos y rúbricas de los talleres son propuestas didácticas.
''',encoding='utf-8')
    (WORK/'Mapa_de_cobertura.json').write_text(json.dumps(all_mapping,ensure_ascii=False,indent=2))
    print('Índice HTML, Word para exportar y guía preparados.',flush=True)

def finish():
    offices=json.loads((WORK/'Verificacion_Office.json').read_text())
    browsers=json.loads((WORK/'Verificacion_navegador.json').read_text())
    assert len(offices)==len(browsers)==8
    assert sum(r['questions'] for r in browsers)==218
    assert sum(r['slides'] for r in offices)==286
    pdf=fitz.open(OUT/'00_Indice.pdf')
    pdftext=' '.join(page.get_text() for page in pdf)
    assert all(t['title'].replace(' ','') in ''.join(pdftext.split()) for u in UNITS for t in u['topics'])
    fields=[]
    for u in UNITS:
        files=[p for p in Path(u['folder']).iterdir() if p.is_file()]
        assert {p.name for p in files}==set(u['files'].values()),u['id']
        verify=json.loads((WORK/u['id']/'Verificacion_medios.json').read_text())
        assert all(verify[k]['fully_decoded_without_errors'] for k in ('video','audio'))
        for p in files:
            assert u['id'] in p.name and slug(u['title']) in p.name and u['topics'][0]['num'] in p.name and u['topics'][-1]['num'] in p.name
            fields.append(p)
    assert len(fields)==40
    for href in re.findall(r'href="([^"]+)"',(OUT/'00_Indice.html').read_text()):
        assert (OUT/unquote(html.unescape(href))).exists(),href
    report={'materias':2,'unidades':8,'secciones':27,'archivos_principales':40,'diapositivas':286,
            'preguntas_practicas_locales':218,'workshops':offices,'browser':browsers,
            'educaplay_verificacion_remota':'Bloqueada por política de red. Perfiles asignados según el usuario; enlaces de juegos no verificados.',
            'pdf_indice_paginas':len(pdf),'max_longitud_ruta_zip':max(len((Path('Entrega')/p.relative_to(OUT)).as_posix()) for p in OUT.rglob('*') if p.is_file())}
    (WORK/'Verificacion_final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    hashes=[hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(OUT).as_posix() for p in sorted(OUT.rglob('*')) if p.is_file()]
    (WORK/'MANIFIESTO_SHA256.txt').write_text('\n'.join(hashes)+'\n')
    with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file():z.write(p,(Path('Entrega')/p.relative_to(OUT)).as_posix())
    with zipfile.ZipFile(ZIP) as z:
        assert z.testzip() is None
        count=len(z.namelist())
        assert count==44,count
        assert all(sum(n.endswith(ext) for n in z.namelist())==8 for ext in ('.pptx','.mp4','.mp3','.docx'))
    print(json.dumps({'zip':str(ZIP),'bytes':ZIP.stat().st_size,'archivos':count,'unidades':8,'cinco_archivos_principales_por_unidad':True,'integridad':'OK','educaplay_juegos_remotos_verificados':False},ensure_ascii=False,indent=2),flush=True)

def slug(s):
    import unicodedata
    return re.sub(r'[^A-Za-z0-9]+','_',unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode()).strip('_')

if __name__=='__main__':
    if '--final' in sys.argv:finish()
    else:prepare()
