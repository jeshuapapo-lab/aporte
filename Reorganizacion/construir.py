from pathlib import Path
import copy, csv, html, json, re, shutil, subprocess, sys, unicodedata
from concurrent.futures import ProcessPoolExecutor, as_completed
from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from docx import Document
from docx.shared import Inches, Pt, RGBColor

BASE = Path('/workspace')
SOURCE = BASE / 'Proyecto_Educativo_Completo'
WORK = BASE / 'reorganizacion-unidades'
OUT = BASE / 'Entrega'
sys.path.insert(0, '/workspace/edu-build')
from contenido import SYLLABI, SLIDES
from generar_documentos import Canvas

def slug(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()).strip('_')

def run(args):
    proc = subprocess.run([str(s) for s in args], capture_output=True, text=True)
    if proc.returncode:
        raise RuntimeError(proc.stderr[-5000:])
    return proc

def probe(path):
    return json.loads(run(['ffprobe', '-v', 'error', '-show_format', '-show_streams', '-show_chapters', '-of', 'json', path]).stdout)

RECORDS = json.loads((SOURCE / 'Validacion.json').read_text())['records']
BY_ID = {r['id']: r for r in RECORDS}
with (BASE / 'Educaplay_6_cuentas/Control_Educaplay_6_cuentas.csv').open(encoding='utf-8-sig', newline='') as f:
    PROFILES = {r['ID']: r for r in csv.DictReader(f)}

def topic(code, number, title):
    ident = code + '_' + number.replace('.', '_')
    r = BY_ID[ident]
    src = SOURCE / r['path']
    if ident != 'NE_1_1':
        content = json.loads((BASE / 'edu-full/content' / (ident + '.json')).read_text())
    else:
        content = dict(objective='Proponer y justificar una página de inicio que comunique la oferta, jerarquice recorridos y pueda evaluarse.',
                       case='Minga Cacao es un negocio ficticio de Guayaquil. Su inicio solo dice «Bienvenidos», contiene seis botones iguales con la etiqueta «Más» y no aclara qué vende. No existen datos confirmados de plazos de entrega, certificaciones o testimonios. En un mes hipotético hubo 2.000 sesiones que vieron el inicio y 50 completaron una solicitud de información, una vez por sesión.',
                       tasks=['Identifiquen tres dificultades para una persona que busca un regalo y justifiquen un cambio para cada una.',
                              'Redacten un título de hasta 14 palabras, apoyo de hasta 35 palabras y una acción prioritaria con etiqueta de 2–5 palabras y destino explícito.',
                              'Construyan un boceto móvil con identidad, navegación, propuesta, acción principal, apoyo y pie; expliquen dos recorridos de visitantes y dos comprobaciones de accesibilidad.',
                              'Calculen la proporción de sesiones que completan la solicitud, indicando fórmula y resultado. Expliquen por qué no equivale a ventas y cómo evaluarían el rediseño sin atribuir causalidad sin evidencia.'],
                       source='Sílabo UPS, Negocios Electrónicos, 1.1, p. 1; W3C, WCAG 2.2 (2023).')
    with (src / 'Recursos/Banco_preguntas.csv').open(encoding='utf-8-sig', newline='') as f:
        bank = list(csv.DictReader(f))
    questions = []
    for row in bank:
        lowered = {k.lower(): v for k, v in row.items()}
        questions.append(dict(section=number, title=title, q=lowered['pregunta'],
                              options=[row[c] for c in 'ABCD'], answer='ABCD'.index(lowered['correcta']),
                              feedback=lowered['retroalimentacion']))
    assert len(questions) == (10 if ident == 'NE_1_1' else 8)
    result=dict(id=ident, num=number, title=title, source=str(src), content=content,
                questions=questions, profile=PROFILES[ident]['Perfil_aportado_por_usuario'])
    direct=WORK/'Enlaces_directos_verificados_documentalmente.json'
    if direct.exists():
        evidence=json.loads(direct.read_text())[ident]
        result.update(activity_url=evidence['url'],activity_evidence=evidence)
    return result

def catalog():
    units = []
    for m in SYLLABI:
        code = 'NE' if m['materia'] == 'Negocios Electrónicos' else 'PND'
        for n, (title, topics) in enumerate(m['unidades'], 1):
            folder = OUT / slug(m['materia']) / f'Unidad_{n}_{slug(title)}'
            span = f'{topics[0][0]}_a_{topics[-1][0]}'
            stem = f'{code}_U{n}_{slug(title)}__{span}'
            unit = dict(id=f'{code}_U{n}', code=code, number=n, materia=m['materia'], title=title,
                        folder=str(folder), stem=stem, topics=[topic(code, num, name) for num, name in topics])
            unit['files'] = {kind: f'{stem}__{suffix}' for kind, suffix in {
                'pptx':'Presentacion.pptx', 'video':'Video.mp4', 'audio':'Audio.mp3',
                'workshop':'Taller_grupal.docx', 'educaplay':'Educaplay.html'}.items()}
            units.append(unit)
    return units

def merge_pptx(unit):
    dest = Path(unit['folder'])
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    for kind in ('cover', 'agenda'):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        c = Canvas(slide)
        c.rect(0, 0, 1600, 900, '#112B3C')
        c.rect(0, 0, 1600, 12, '#087E8B')
        c.text(80, 55, 1420, unit['materia'].upper(), 26, '#E7AE53', True, max_h=80)
        if kind == 'cover':
            c.text(80, 180, 1430, f"Unidad {unit['number']}", 38, '#D6E4E8', True)
            c.text(80, 265, 1420, unit['title'], 58, '#FFFFFF', True, max_h=300)
            c.text(85, 650, 1400, 'Secciones ' + ' · '.join(t['num'] for t in unit['topics']), 32, '#D6E4E8', max_h=120)
        else:
            c.text(80, 165, 1420, 'Secciones y contenidos de la unidad', 45, '#FFFFFF', True, max_h=110)
            for i, t in enumerate(unit['topics']):
                c.text(90, 310 + i * 125, 1410, f"{t['num']} · {t['title']}", 32, '#FFFFFF', max_h=115)
        slide.notes_slide.notes_text_frame.text = f"{unit['materia']}. Unidad {unit['number']}: {unit['title']}. Se conservan íntegros los contenidos de las secciones del sílabo."
    section_slides = []
    for t in unit['topics']:
        donor = Presentation(Path(t['source']) / 'Presentacion.pptx')
        visited = set()
        def rename(part):
            if id(part) in visited:
                return
            visited.add(id(part))
            old = Path(str(part.partname))
            part._partname = PackURI(str(old.parent / (t['id'] + '_' + old.name)))
            for rel in part.rels.values():
                if not rel.is_external:
                    rename(rel.target_part)
        first = len(prs.slides) + 1
        for src_slide in donor.slides:
            target = prs.slides.add_slide(prs.slide_layouts[6])
            target._element.cSld.set('name', f"{unit['id']} · Sección {t['num']} · {t['title']}")
            bg = src_slide._element.cSld.find(qn('p:bg'))
            if bg is not None:
                target._element.cSld.insert(0, copy.deepcopy(bg))
            relmap = {}
            for rel in src_slide.part.rels.values():
                if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
                    continue
                if rel.is_external:
                    rid = target.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    rename(rel.target_part)
                    rid = target.part.rels.get_or_add(rel.reltype, rel.target_part)
                relmap[rel.rId] = rid
            for shape in src_slide.shapes:
                element = copy.deepcopy(shape.element)
                for node in element.iter():
                    for attr, value in list(node.attrib.items()):
                        if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in relmap:
                            node.set(attr, relmap[value])
                target.shapes._spTree.insert_element_before(element, 'p:extLst')
            if src_slide.has_notes_slide:
                target.notes_slide.notes_text_frame.text = (
                    f"{unit['materia']} · Unidad {unit['number']}: {unit['title']}\nSección {t['num']}: {t['title']}\n\n"
                    + src_slide.notes_slide.notes_text_frame.text)
        section_slides.append(dict(num=t['num'], title=t['title'], first=first, last=len(prs.slides)))
    prs.core_properties.title = f"Unidad {unit['number']}. {unit['title']}"
    prs.core_properties.subject = unit['materia'] + ' · ' + ' · '.join(t['num'] + ' ' + t['title'] for t in unit['topics'])
    prs.core_properties.author = 'Proyecto educativo · elaboración asistida por IA'
    path = dest / unit['files']['pptx']
    prs.save(path)
    check = Presentation(path)
    assert len(check.slides) == 2 + 10 * len(unit['topics'])
    unit['slides'] = section_slides
    unit['slide_count'] = len(check.slides)

def make_workshop(unit):
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(.7)
        s.left_margin = s.right_margin = Inches(.75)
        s.header.paragraphs[0].text = f"{unit['materia']} · Unidad {unit['number']}"
        s.footer.paragraphs[0].text = 'Taller colaborativo · Caso y datos ficticios · Propuesta didáctica'
    d.styles['Normal'].font.name = 'Calibri'
    d.styles['Normal'].font.size = Pt(10.5)
    d.styles['Normal'].paragraph_format.space_after = Pt(5)
    d.add_heading('Taller grupal', 0)
    d.add_heading(f"Unidad {unit['number']}: {unit['title']}", 1)
    d.add_paragraph('Equipo: ____________________  Curso: __________  Fecha: __________')
    d.add_heading('Objetivo y alcance', 1)
    d.add_paragraph('Resolver colaborativamente las situaciones de las siguientes secciones, justificar las decisiones con evidencia y relacionar los resultados en una propuesta de la unidad:')
    for t in unit['topics']:
        d.add_paragraph(f"{t['num']} · {t['title']}. {t['content']['objective']}", style='List Bullet')
    d.add_heading('Organización del equipo', 1)
    d.add_paragraph('Equipos de 3 a 5 integrantes. Duración sugerida: 100–140 minutos, ajustable por el docente. Estos tamaños, tiempos y la rúbrica son propuestas didácticas, no disposiciones atribuidas al profesor.')
    table = d.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    for cell, text in zip(table.rows[0].cells, ['Rol', 'Responsabilidad', 'Integrante']): cell.text = text
    for role, duty in [('Coordinación','Organiza tiempos y asegura que todos discutan cada sección.'),
                       ('Análisis','Contrasta los conceptos y documenta los argumentos del equipo.'),
                       ('Diseño y registro','Integra diagramas, tablas y respuestas en un único documento.'),
                       ('Validación','Comprueba cálculos, fuentes, consistencia y supuestos.'),
                       ('Vocería','Expone acuerdos y responde preguntas con apoyo del equipo.')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = role, duty, '________________'
    d.add_paragraph('En equipos de tres, combinar diseño con coordinación y vocería con análisis. Rotar la función de validación entre secciones. Cada respuesta debe ser revisada por un integrante distinto de quien la redactó.')
    d.add_heading('Secuencia y producto común', 1)
    for line in ['Revisar juntos la presentación y los objetivos; acordar cómo demostrarán cada resultado (10 minutos).',
                 'Resolver las estaciones de todas las secciones. Discutir y registrar los desacuerdos y su resolución; no limitarse a reunir fragmentos sin revisión común (60–90 minutos).',
                 'Validar conceptos, cálculos y diagramas con revisión cruzada, y preparar la síntesis de la unidad (20 minutos).',
                 'Presentar una defensa grupal de cinco minutos; cualquier integrante debe poder explicar una decisión de una sección distinta a la que redactó (10–20 minutos).']:
        d.add_paragraph(line, style='List Number')
    d.add_paragraph('Entrega: un Word o PDF del equipo, con respuestas identificadas por sección y ejercicio, evidencias gráficas o tablas solicitadas, síntesis final, fuentes y registro de contribuciones. Los diagramas pueden insertarse como imágenes legibles y adjuntarse en formato editable cuando corresponda.')
    for t in unit['topics']:
        d.add_heading(f"Sección {t['num']} · {t['title']}", 1)
        d.add_heading('Situación de análisis', 2)
        d.add_paragraph(t['content']['case'])
        if t['id']=='PND_4_4':
            dataset=Path(t['source'])/'Recursos/Casos_soporte_simulados.csv'
            resources=Path(unit['folder'])/'Recursos'
            resources.mkdir(exist_ok=True)
            shutil.copy2(dataset,resources/'PND_U4_4.4_Estudio_de_caso_Datos_simulados.csv')
            with dataset.open(encoding='utf-8-sig',newline='') as f:
                records=list(csv.DictReader(f))
            d.add_paragraph('Datos para resolver el taller: cada escenario contiene 100 casos. El cumplimiento se define como respuesta útil en un máximo de 48 horas; cada caso requiere 4 horas de trabajo y el resto es espera. Una corrección se cuenta una sola vez por caso. El CSV de 200 filas está en la subcarpeta Recursos de esta unidad. El cuadro siguiente permite realizar los cálculos sin abrir otro archivo.')
            data=d.add_table(rows=1,cols=5);data.style='Light Shading Accent 1'
            for cell,text in zip(data.rows[0].cells,['Escenario','Cantidad de casos','Ciclo por caso (h)','Casos con corrección','Cumplen 48 h']):cell.text=text
            for period in ('inicial','futuro_simulado'):
                times=sorted({int(r['horas_ciclo']) for r in records if r['periodo']==period})
                for hours in times:
                    group=[r for r in records if r['periodo']==period and int(r['horas_ciclo'])==hours]
                    vals=['Inicial' if period=='inicial' else 'Futuro simulado',str(len(group)),str(hours),str(sum(int(r['requiere_correccion']) for r in group)),str(sum(int(r['cumple_48_horas']) for r in group))]
                    for cell,text in zip(data.add_row().cells,vals):cell.text=text
            assert len(records)==200
            assert len(data.rows)==5 and sum(int(row.cells[1].text) for row in data.rows[1:])==200
        for i, task in enumerate(t['content']['tasks'], 1):
            # Adaptación explícita al trabajo grupal, sin conservar instrucciones de entrega individual.
            d.add_heading(f"Ejercicio {t['num']}.{i}", 2)
            plurals = {'Asigna':'Asignen','Asocia':'Asocien','Audita':'Auditen','Calcula':'Calculen',
                       'Clasifica':'Clasifiquen','Compara':'Comparen','Construye':'Construyan','Corrige':'Corrijan',
                       'Define':'Definan','Delimita':'Delimiten','Describe':'Describan','Dibuja':'Dibujen',
                       'Diseña':'Diseñen','Elabora':'Elaboren','Especifica':'Especifiquen','Evalúa':'Evalúen',
                       'Explica':'Expliquen','Formula':'Formulen','Identifica':'Identifiquen','Justifica':'Justifiquen',
                       'Modela':'Modelen','Modifica':'Modifiquen','Propón':'Propongan','Recalcula':'Recalculen',
                       'Recorre':'Recorran','Redacta':'Redacten','Relaciona':'Relacionen','Representa':'Representen'}
            if t['id']=='PND_4_4' and i==3:
                task='Calcula el promedio de ciclo, el porcentaje de casos con corrección y el cumplimiento de 48 horas para ambos escenarios, usando el cuadro anterior o el CSV. Muestra fórmulas, denominadores y variaciones; explica qué ocurre con los cinco casos futuros de cien horas.'
            first = task.split(' ', 1)[0]
            if first in plurals:
                task = plurals[first] + task[len(first):]
            for singular, plural in [('exigirías','exigirían'),('evaluarías','evaluarían'),('usarías','usarían'),
                                     ('medirías','medirían'),('comprobarías','comprobarían')]:
                task=task.replace(singular,plural)
            d.add_paragraph(task)
            d.add_paragraph('Respuesta y evidencia acordadas: ____________________________________________\n_____________________________________________________________________')
        d.add_paragraph('Revisó esta sección: ____________________  Ajuste acordado tras la revisión: ____________________')
    d.add_heading('Síntesis y defensa de la unidad', 1)
    for task in ['Explicar dos relaciones concretas entre las secciones: qué decisión de una afecta a otra y por qué.',
                 'Elegir una decisión prioritaria y justificarla mediante una evidencia o un cálculo trabajado. Delimitar qué no permiten concluir los datos.',
                 'Registrar un desacuerdo real del equipo, las alternativas consideradas y el criterio usado para resolverlo.',
                 'Proponer cómo comprobarían la solución en otro caso y qué información faltaría antes de aplicarla.']:
        d.add_paragraph(task, style='List Number')
    d.add_heading('Rúbrica grupal propuesta · 20 puntos', 1)
    rub = d.add_table(rows=1, cols=3)
    rub.style = 'Light Shading Accent 1'
    for c, text in zip(rub.rows[0].cells, ['Criterio', 'Evidencia esperada', 'Puntos']): c.text = text
    for a,b,c in [('Cobertura y corrección','Respuestas de todas las secciones con conceptos, cálculos y diagramas correctos.','6'),
                  ('Argumentación e integración','Decisiones justificadas, dos relaciones entre secciones y límites de la evidencia.','5'),
                  ('Colaboración verificable','Roles, revisión cruzada, contribuciones y explicación del consenso.','4'),
                  ('Comunicación y defensa','Documento identificado, evidencias legibles, fuentes y defensa compartida.','5')]:
        cells = rub.add_row().cells
        for cell, value in zip(cells, (a,b,c)): cell.text = value
    d.add_paragraph('Por criterio: completo y justificado = 100 %; correcto con omisión relevante = 75 %; parcial = 50 %; mínimo = 25 %; ausente = 0 %. El docente puede ajustar esta escala y valorar individualmente la defensa.')
    d.add_heading('Contribuciones y uso de fuentes', 1)
    contrib = d.add_table(rows=1, cols=4)
    contrib.style = 'Light Shading Accent 1'
    for c,text in zip(contrib.rows[0].cells, ['Integrante','Aporte concreto','Evidencia/sección','Revisión de un compañero']): c.text=text
    for _ in range(5):
        for c in contrib.add_row().cells: c.text='________________'
    d.add_paragraph('Fuentes consultadas y herramientas de IA utilizadas: ______________________________\nQué sugerencias revisó o corrigió el equipo: _______________________________________')
    d.add_paragraph('Base curricular: sílabo UPS 2026–2027, ' + unit['materia'] + f", Unidad {unit['number']}. " + ' '.join(t['content']['source'] for t in unit['topics']))
    path = Path(unit['folder']) / unit['files']['workshop']
    d.save(path)
    check = Document(path)
    alltext = '\n'.join(p.text for p in check.paragraphs)
    assert 'Trabajo individual' not in alltext
    assert all(f"Sección {t['num']} · {t['title']}" in alltext for t in unit['topics'])
    assert len(contrib.rows) == 6

def make_quiz(unit):
    data = [q for t in unit['topics'] for q in t['questions']]
    rows = ''.join(f'<tr><td>{html.escape(t["num"])}</td><td>{html.escape(t["title"])}</td><td><a href="{t["profile"]}" target="_blank" rel="noopener">Abrir perfil · buscar {t["num"]}</a></td></tr>' for t in unit['topics'])
    page = '''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Educaplay · __TITLE__</title><style>*{box-sizing:border-box}body{max-width:1000px;margin:32px auto;padding:0 20px;font:17px/1.6 system-ui;color:#18394a;background:#f5f2eb}h1{line-height:1.2}a{color:#087078;overflow-wrap:anywhere}section{background:white;border-radius:14px;padding:24px;margin:24px 0}table{width:100%;border-collapse:collapse}td,th{text-align:left;border-bottom:1px solid #ccd9d9;padding:12px;overflow-wrap:anywhere}button{font:inherit;padding:12px 16px;border:2px solid #087e8b;border-radius:8px;background:white;color:#18394a;cursor:pointer;margin:7px 0;text-align:left}#options button{display:block;width:100%}button:focus-visible,a:focus-visible,summary:focus-visible{outline:4px solid #b9432e;outline-offset:3px}button:disabled{cursor:default;color:#344d58}button.correct{background:#d3eee1}button.wrong{background:#f9dbd4}#feedback{padding:12px;border-left:5px solid #087e8b;background:#edf6f5}#section{color:#087078;font-weight:bold}#next,#restart{background:#087e8b;color:white}#result{font-weight:bold;font-size:1.3rem}@media(max-width:600px){body{padding:0 12px}section{padding:16px}table{font-size:14px}td,th{padding:7px}}</style></head><body><h1>Unidad __N__: __TITLE__</h1><p>__SUBJECT__ · secciones __SPAN__</p><section><h2>Acceso a Educaplay por sección</h2><p>Estos son los perfiles facilitados por el estudiante. Abre el perfil indicado y selecciona el juego del tema correspondiente. Los enlaces individuales a los juegos no se recibieron y su correspondencia no se ha podido comprobar desde este entorno.</p><table><thead><tr><th>Sección</th><th>Tema</th><th>Acceso recibido</th></tr></thead><tbody>__ROWS__</tbody></table></section><section><h2>Práctica complementaria de la unidad</h2><p>Cuestionario local con todas las preguntas preparadas de las secciones. Esta práctica funciona en el navegador y es independiente del sitio Educaplay. Elige una respuesta y revisa la explicación. Sin límite de tiempo ni descuento por error.</p><p id="counter"></p><p id="section"></p><h3 id="question"></h3><div id="options"></div><p id="feedback" role="status" aria-live="polite" hidden></p><button id="next" hidden>Siguiente pregunta</button><p id="result" role="status" aria-live="polite"></p><button id="restart">Reiniciar práctica</button></section><script id="bank" type="application/json">__DATA__</script><script>const questions=JSON.parse(document.getElementById('bank').textContent);let pos=0,score=0,locked=false;const el=id=>document.getElementById(id);function render(){locked=false;el('feedback').hidden=true;el('next').hidden=true;el('result').textContent='';el('counter').textContent='Pregunta '+(pos+1)+' de '+questions.length;const q=questions[pos];el('section').textContent='Sección '+q.section+' · '+q.title;el('question').textContent=q.q;el('options').replaceChildren();q.options.forEach((text,i)=>{const b=document.createElement('button');b.type='button';b.textContent=String.fromCharCode(65+i)+') '+text;b.addEventListener('click',()=>answer(i));el('options').append(b)});}function answer(i){if(locked)return;locked=true;const q=questions[pos];if(i===q.answer)score++;[...el('options').children].forEach((b,j)=>{b.disabled=true;if(j===q.answer)b.classList.add('correct');else if(j===i)b.classList.add('wrong')});el('feedback').textContent=(i===q.answer?'Correcto. ':'Respuesta incorrecta. La correcta es '+String.fromCharCode(65+q.answer)+'. ')+q.feedback;el('feedback').hidden=false;el('next').hidden=false;el('next').textContent=pos===questions.length-1?'Ver resultado':'Siguiente pregunta';}el('next').onclick=()=>{if(!locked)return;if(pos<questions.length-1){pos++;render();el('question').setAttribute('tabindex','-1');el('question').focus();}else{el('result').textContent='Resultado: '+score+'/'+questions.length;el('next').hidden=true;}};el('restart').onclick=()=>{pos=0;score=0;render()};render();</script></body></html>'''
    for key,value in {'__TITLE__':html.escape(unit['title']), '__N__':str(unit['number']),
                      '__SUBJECT__':html.escape(unit['materia']), '__SPAN__':', '.join(t['num'] for t in unit['topics']),
                      '__ROWS__':rows, '__DATA__':json.dumps(data,ensure_ascii=False).replace('</','<\\/')}.items():
        page=page.replace(key,value)
    if all(t.get('activity_url') for t in unit['topics']):
        for t in unit['topics']:
            page=page.replace('href="'+t['profile']+'" target="_blank" rel="noopener">Abrir perfil · buscar '+t['num'],
                              'href="'+t['activity_url']+'" target="_blank" rel="noopener">Abrir juego '+t['num'])
        page=page.replace('Estos son los perfiles facilitados por el estudiante. Abre el perfil indicado y selecciona el juego del tema correspondiente. Los enlaces individuales a los juegos no se recibieron y su correspondencia no se ha podido comprobar desde este entorno.',
                          'Cada sección enlaza directamente a su actividad de Educaplay. Los 27 destinos se contrastaron con los números, títulos y preguntas del PDF aportado por el estudiante. Ese informe registra comprobación pública de página y reproductor el 30/10/2026; no se repitió la comprobación remota desde este entorno.')
    (Path(unit['folder']) / unit['files']['educaplay']).write_text(page,encoding='utf-8')
    unit['questions']=len(data)

def media(unit):
    work=WORK/unit['id'];work.mkdir(exist_ok=True)
    sources=[]
    for t in unit['topics']:
        path=Path(t['source'])/'Video.mp4'
        info=probe(path)
        audio=next(s for s in info['streams'] if s['codec_type']=='audio')
        if audio['sample_rate']!='48000':
            normalized=work/(t['id']+'_48000.mp4')
            if not normalized.exists():
                run(['ffmpeg','-v','warning','-y','-i',path,'-map','0:v:0','-map','0:a:0','-c:v','copy','-c:a','aac','-b:a','128k','-ar','48000','-ac','1','-video_track_timescale','12288',normalized])
            path=normalized
        sources.append(path)
    listing=work/'video_concat.txt'
    listing.write_text(''.join("file '"+str(p).replace("'","'\\''")+"'\n" for p in sources))
    def meta(paths,name):
        offset=0;chapters=[]
        lines=[';FFMETADATA1','title='+unit['title'],'artist=Proyecto educativo - '+unit['materia'],'comment=Unidad '+str(unit['number'])+'; secciones '+', '.join(t['num'] for t in unit['topics'])]
        for t,path in zip(unit['topics'],paths):
            duration=round(float(probe(path)['format']['duration'])*1000)
            title=(t['num']+' - '+t['title']).replace('=','\\=').replace(';','\\;').replace('#','\\#')
            lines += ['[CHAPTER]','TIMEBASE=1/1000','START='+str(offset),'END='+str(offset+duration),'title='+title]
            chapters.append(dict(section=t['num'],title=t['title'],start=offset/1000,end=(offset+duration)/1000))
            offset+=duration
        target=work/name;target.write_text('\n'.join(lines)+'\n')
        return target,chapters
    metadata,chapters=meta(sources,'video.ffmeta')
    video=Path(unit['folder'])/unit['files']['video']
    if not video.exists():
        run(['ffmpeg','-v','warning','-y','-f','concat','-safe','0','-i',listing,'-f','ffmetadata','-i',metadata,
             '-map','0:v:0','-map','0:a:0','-map_metadata','1','-map_chapters','1','-c:v','copy',
             '-c:a','aac','-b:a','128k','-ar','48000','-ac','1','-af','aresample=async=1:first_pts=0',
             '-movflags','+faststart','-threads','2',video])
    audios=[Path(t['source'])/'Audio.mp3' for t in unit['topics']]
    audio_meta,audio_chapters=meta(audios,'audio.ffmeta')
    args=['ffmpeg','-v','warning','-y']
    for p in audios:args += ['-i',p]
    args += ['-f','ffmetadata','-i',audio_meta]
    filt=';'.join(f'[{i}:a:0]aresample=44100,asetpts=PTS-STARTPTS[a{i}]' for i in range(len(audios)))+';'+''.join(f'[a{i}]' for i in range(len(audios)))+f'concat=n={len(audios)}:v=0:a=1[out]'
    audio=Path(unit['folder'])/unit['files']['audio']
    if not audio.exists():
        run(args+['-filter_complex',filt,'-map','[out]','-map_metadata',str(len(audios)),'-map_chapters',str(len(audios)),
                  '-c:a','libmp3lame','-b:a','96k','-ar','44100','-ac','1','-id3v2_version','3','-threads','2',audio])
    result=dict(id=unit['id'],video=probe(video),audio=probe(audio),chapters=chapters,audio_chapters=audio_chapters)
    (work/'media.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    return result

def main():
    units=catalog();OUT.mkdir(exist_ok=True)
    for u in units:
        Path(u['folder']).mkdir(parents=True,exist_ok=True)
        merge_pptx(u);make_workshop(u);make_quiz(u)
        assert len([p for p in Path(u['folder']).iterdir() if p.is_file()]) in (3,5)
        print('Documentos y actividad:',u['id'],u['slide_count'],'diapositivas',u['questions'],'preguntas',flush=True)
    (WORK/'catalogo.json').write_text(json.dumps(units,ensure_ascii=False,indent=2))
    with ProcessPoolExecutor(max_workers=2) as pool:
        futures={pool.submit(media,u):u for u in units}
        for future in as_completed(futures):
            result=future.result();u=futures[future]
            print('Medios:',u['id'],'video',result['video']['format']['duration'],'audio',result['audio']['format']['duration'],flush=True)
    (WORK/'GENERADO').write_text('8 unidades · 40 archivos principales · 27 secciones')

if __name__=='__main__':main()
