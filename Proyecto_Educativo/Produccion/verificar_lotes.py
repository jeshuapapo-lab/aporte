import json,time,subprocess,zipfile,sys,os
from pathlib import Path
import fitz
from pptx import Presentation
from docx import Document
ROOT=Path('/workspace/Proyecto_Educativo_Completo');BASE=Path('/workspace/edu-full')
def run(cmd):
 r=subprocess.run(cmd,capture_output=True,text=True)
 if r.returncode:raise RuntimeError(r.stderr[-2500:])
 return r
def check(folder):
 for name in ['Presentacion.pptx','Video.mp4','Audio.mp3','Tarea.docx','Actividad_Educaplay.txt','Actividad_Interactiva.html']:
  assert (folder/name).stat().st_size>100,name
 for name in ['Presentacion.pptx','Tarea.docx']:
  with zipfile.ZipFile(folder/name) as z:assert z.testzip() is None
 prs=Presentation(folder/'Presentacion.pptx');assert len(prs.slides)==10
 assert all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
 assert len(Document(folder/'Tarea.docx').paragraphs)>10
 run(['soffice','-env:UserInstallation=file:///tmp/edu-full-lo','--headless','--convert-to','pdf','--outdir',str(folder),str(folder/'Presentacion.pptx'),str(folder/'Tarea.docx')])
 pages={}
 for name in ['Presentacion.pdf','Tarea.pdf']:
  d=fitz.open(folder/name);pages[name]=len(d)
  assert all(len(pg.get_text().strip())>80 for pg in d),name
  if name=='Presentacion.pdf':
   assert len(d)==10
   for pg in d:
    for b in pg.get_text('dict')['blocks']:
     if b['type']==0:
      x0,y0,x1,y1=b['bbox'];assert x0>=-1 and y0>=-1 and x1<=pg.rect.width+1 and y1<=pg.rect.height+1,(name,b)
 results={}
 for name in ['Video.mp4','Audio.mp3']:
  p=folder/name;info=json.loads(run(['ffprobe','-v','error','-show_format','-show_streams','-of','json',str(p)]).stdout)
  results[name]=dict(duration=float(info['format']['duration']),bytes=p.stat().st_size,streams=[{k:s[k] for k in ['codec_type','codec_name','width','height','sample_rate','channels','pix_fmt'] if k in s} for s in info['streams']])
  if name=='Video.mp4':
   assert 120<=results[name]['duration']<=240
   assert info['streams'][0]['codec_name']=='h264'
   assert info['streams'][0]['width']==1920 and info['streams'][0]['height']==1080
  run(['ffmpeg','-v','error','-xerror','-threads','1','-i',str(p),'-f','null','-'])
  results[name]['decoded_entirely']=True
 record=dict(folder=str(folder.relative_to(ROOT)),office_reopened=True,pdf_pages=pages,media=results)
 (folder/'VERIFICACION.json').write_text(json.dumps(record,ensure_ascii=False,indent=2))
 print('VERIFICADO '+folder.name,flush=True)
 return record
if __name__=='__main__':
 failed=[]
 while True:
  for marker in sorted(ROOT.rglob('VALIDADO_PRODUCCION.json')):
   if (marker.parent/'VERIFICACION.json').exists():continue
   if (marker.parent/'ERROR_VERIFICACION.json').exists():continue
   try:check(marker.parent)
   except Exception as e:
    failed.append(dict(folder=str(marker.parent),error=str(e)));print('ERROR VALIDACIÓN '+str(marker.parent)+' '+str(e),flush=True)
    (marker.parent/'ERROR_VERIFICACION.json').write_text(json.dumps(failed[-1],ensure_ascii=False))
  (BASE/'estado_validacion.json').write_text(json.dumps(dict(validated=len(list(ROOT.rglob('VERIFICACION.json'))),failed=failed),ensure_ascii=False,indent=2))
  state=json.loads((BASE/'estado.json').read_text())
  if (BASE/'CATALOGO_COMPLETO').exists() and state['pending']==0 and (len(list(ROOT.rglob('VERIFICACION.json')))==26 or state['failed']):
   break
  time.sleep(5)
