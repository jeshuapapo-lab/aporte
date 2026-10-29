import csv, hashlib, importlib.metadata as metadata, json, shutil, subprocess, sys, zipfile
from pathlib import Path
import fitz
from pptx import Presentation
from docx import Document
from contenido import *

README='''# Proyecto educativo · piloto y planificación

## Qué contiene esta entrega

Se analizaron los dos sílabos de la Universidad Politécnica Salesiana, período 2026–2027: **Negocios Electrónicos** (12 subtemas) y **Procesos para Negocios Digitales** (15 subtemas). La matriz conserva las cuatro unidades de cada materia y los 27 subtemas en el orden de sus documentos.

**Alcance desarrollado: un piloto propuesto, Negocios Electrónicos, unidad 1, tema 1.1: «Las páginas de inicio».** Los PDF no asignan temas a un estudiante concreto. Los otros 26 subtemas solo están planificados; no se presentan como producidos ni como asignados. Desarrollarlos todos supondría 135 entregables principales.

## Archivos del piloto

Carpeta: `Negocios_Electronicos/Tema_1_1_Las_paginas_de_inicio/`

| Archivo | Contenido y estado |
|---|---|
| Presentacion.pptx | 10 diapositivas 16:9 editables, con explicación, diagrama de página, caso, aplicación, conclusión y referencias. Incluye guion en las notas. Abierta con LibreOffice y reabierta con python-pptx. |
| Video.mp4 | 3:38,125 · Full HD 1920×1080 · 24 fps · H.264/AAC. Narración sintética neuronal en español con fonetización latinoamericana es-419, ilustraciones y diagramas originales, subtítulos incrustados, transiciones y barra animada. Música instrumental original de teclado, a bajo nivel y atenuada durante la voz. |
| Audio.mp3 | Aproximadamente 3:17 · MP3 mono 44,1 kHz. Explicación independiente con contexto y pregunta de reflexión; no requiere mirar la presentación. Voz neuronal es-419, sin música para favorecer la escucha. |
| Tarea.docx | 4 páginas editables: objetivo, caso, instrucciones, 6 ejercicios, espacio para boceto, comprobaciones y rúbrica propuesta de 20 puntos. |
| Actividad_Educaplay.txt | Cuestionario completo de 10 preguntas con 4 opciones, respuesta correcta, retroalimentación, instrucciones y pasos para publicar. **Preparado, no publicado. No hay enlace de una actividad alojada.** |

Extras: actividad HTML interactiva sin conexión, banco CSV auxiliar, subtítulos SRT, guiones, guía docente con orientaciones de corrección, fotogramas reales y versiones PDF de la presentación y la tarea.

## Cómo usar los archivos

1. Extrae todo el ZIP antes de abrirlo. No necesitas ejecutar código para usar los materiales.
2. Abre PPTX y DOCX en PowerPoint/Word o una suite compatible. Los PDF sirven de vista previa y no sustituyen los editables.
3. Reproduce MP4 y MP3 en un reproductor que admita H.264/AAC y MP3. Ambos se reprodujeron en Chromium y se decodificaron por completo con FFmpeg.
4. Para practicar, abre `Actividad_Interactiva.html` en un navegador. Si tu navegador administrado bloquea archivos locales, sirve únicamente la carpeta del tema con `python -m http.server 8001 --bind 127.0.0.1 --directory RUTA_DEL_TEMA` y accede localmente. Es una práctica independiente, no una actividad de Educaplay; no recopila ni transmite respuestas.
5. Para publicar en Educaplay usa tu cuenta autorizada y sigue `Actividad_Educaplay.txt`. Prueba el enlace real con permisos de estudiante antes de compartirlo. El CSV es auxiliar: no se afirma compatibilidad con una importación automática del servicio.
6. La guía docente y el banco contienen respuestas. Comparte con estudiantes solo los archivos que correspondan a tu modalidad de evaluación.

## Alineación didáctica del piloto

Base curricular: U1 «Estructura del negocio electrónico», resultado «Conoce la estructura de un negocio electrónico», subtema 1.1 «Las páginas de inicio», p. 1 del sílabo.

| Objetivo didáctico propuesto | Desarrollo | Evidencia |
|---|---|---|
| Identificar la función del inicio y sus elementos | Diapositivas 1–3 y explicación multimedia | Tarea 1 y 3; cuestionario 1–2 |
| Redactar una oferta y jerarquizar una acción y recorridos | Diapositivas 4–5 y caso Minga Cacao | Tarea 2–4; cuestionario 3–5 y 10 |
| Justificar comprobaciones de accesibilidad y una medida de resultado | Diapositivas 6–7 | Tarea 5–6; cuestionario 6–9 |
| Aplicar e integrar los criterios en un boceto argumentado | Diapositivas 8–10, conclusión y fuentes | Boceto, justificaciones y revisión con la rúbrica |

Accesibilidad y medición son criterios de desarrollo de este subtema, no nuevos temas atribuidos al sílabo. La rúbrica de 20 puntos, los tiempos de trabajo y la referencia de 8/10 del cuestionario son propuestas propias: no reemplazan disposiciones del profesor. El caso **Minga Cacao** y las cifras usadas son ficticios. No se afirman ventas, certificaciones o resultados reales.

## Planificación de las dos materias

- `Analisis_silabos.md`: tabla completa, resultados de aprendizaje, actividades oficiales, límites y selección pendiente.
- `Matriz_temas.docx`, `.pdf` y `.csv`: listado de todos los temas con los cinco materiales requeridos.
- `Fuentes/`: los dos sílabos PDF originales, incluyendo sus bibliografías y perfil docente. El audio adjuntado por el usuario no se utilizó como fuente curricular ni se incluye en este ZIP.
- `Procesos_para_Negocios_Digitales/SELECCION_PENDIENTE.md`: registra que esta materia está analizada y todavía no tiene subtemas desarrollados.

## Verificación y límites

Consulta `Validacion.md`, `Validacion.json` y `MANIFIESTO_SHA256.txt`. La validación incluye integridad de Office, conversión por LibreOffice, revisión visual de documentos, reproducción en navegador, decodificación integral de medios, códecs y duración, y pruebas de las diez respuestas correctas e incorrectas del interactivo, reinicio, teclado y pantalla de 390 px.

La voz se genera con Kokoro, estilo ef_dora y fonetización es-419. No es una grabación humana. Se comprobó la generación, señal y reproducción; **no se realizó una evaluación perceptiva humana de expresividad o pronunciación**. Revisa auditivamente el piloto antes de entregarlo al profesor. Los subtítulos se segmentaron proporcionalmente dentro de cada frase sintetizada: sus tiempos son orientativos, no una alineación fonética forzada. Las animaciones son transiciones, aparición de subtítulos y progreso; no animación de personajes.

No hubo acceso a una cuenta autorizada de Educaplay, por lo que se entrega la preparación y una práctica local comprobada. No se ha publicado una actividad ni se ha inventado un enlace. Tampoco se ha enviado o publicado este proyecto en GitHub. El checkout de aporte permanece sin cambios.

## Código de producción

`Produccion/` incluye los scripts ejecutados y dependencias para regenerar documentos y multimedia, junto con instrucciones técnicas. El ZIP contiene los archivos finales; **no** incluye el entorno virtual ni los 338 MB de modelo/voces. Su descarga solo es necesaria si deseas sintetizar de nuevo la narración.

## Qué falta para continuar

Confirmar materia y números de los subtemas asignados o seleccionados con autorización del profesor. Con esa información se puede ampliar la producción conservando la estructura y el nivel del piloto. Esta entrega corresponde al análisis y al primer tema de prueba; no es una producción de los 27 temas.
'''

PRODUCTION='''# Regeneración de los materiales

Estos scripts generaron los archivos reales entregados. Requieren Python 3.12, FFmpeg con libx264/libmp3lame/AAC/libass, fuentes Noto Sans y, para las comprobaciones indicadas, Chromium y LibreOffice. Los nombres y rutas de binarios pueden requerir adaptación en otro sistema operativo. No hay llamadas a un servicio de voz remoto: la síntesis funciona localmente una vez descargado el modelo.

1. Crea un entorno virtual e instala `requirements.txt` con pip, conservando TLS y las verificaciones del gestor de paquetes. FFmpeg, Noto Sans, Chromium y LibreOffice son dependencias del sistema, no se distribuyen en el ZIP.
2. Descarga por HTTPS, desde las publicaciones oficiales, `kokoro-v1.0.onnx` y `voices-v1.0.bin`:
   - https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.1/kokoro-v1.0.onnx
   - https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.1/voices-v1.0.bin
   El origen oficial documenta el modelo como Apache 2.0 y el adaptador kokoro-onnx como MIT: https://github.com/thewh1teagle/kokoro-onnx . Los modelos no se redistribuyen en este proyecto.
3. Define `EDU_MODEL_DIR` con la carpeta de esos dos archivos. Para no sobrescribir los entregables o tus ediciones, define `EDU_PROJECT_DIR` con una **carpeta nueva** de salida y `EDU_BUILD_DIR` con una carpeta nueva de trabajo. `EDU_SOURCES_DIR` debe señalar a la carpeta Fuentes de esta entrega. Conserva el proxy y las autoridades de confianza de tu entorno.
4. Ejecuta, en este orden, con el Python del entorno virtual:
   ```sh
   python generar_documentos.py
   python generar_audio.py
   python montar_media.py
   ```
   Los dos primeros scripts escriben documentos, gráficos, guiones y narraciones; el tercero compone música original, normaliza la voz, crea escenas, añade subtítulos y monta el MP4 y el MP3 final. La velocidad efectiva empleada en esta entrega fue 1.0 y la fonetización, es-419. El adaptador de sesión corrige el tipo de entrada `speed` según el esquema float del modelo; no desactiva verificaciones.
5. Opcionalmente ejecuta `validar_interactivo.py` y `validar_multimedia.py` para pruebas reales de navegador. Usan un servidor temporal en loopback y lo detienen al terminar. La ruta del navegador usada en este entorno es `/usr/bin/chromium`.
6. La conversión de vista previa se hizo con LibreOffice, por ejemplo `soffice --headless --convert-to pdf --outdir CARPETA_PDF ARCHIVO.pptx`. Las vistas PDF deben regenerarse tras cambiar los editables.

`empaquetar.py` es el registro específico del empaquetado de esta entrega: espera los PDF convertidos en `EDU_BUILD_DIR/pdf`, fotogramas y el informe acústico de validación. No sustituye los pasos de generación, conversión y prueba. No edites `Validacion.md` para afirmar resultados de una regeneración sin volver a ejecutar las comprobaciones.

La generación reutiliza WAV existentes en su carpeta de trabajo. Si modificas los guiones, usa una carpeta `EDU_BUILD_DIR` nueva o elimina únicamente los WAV/JSON de producción que deban regenerarse; de lo contrario podría reutilizarse narración anterior. No se alteran los sílabos ni el repositorio aporte.
'''

def main():
    office=[]
    for p in ROOT.rglob('*'):
        if p.suffix in ('.pptx','.docx'):
            with zipfile.ZipFile(p) as z:assert z.testzip() is None
            if p.suffix=='.pptx':
                prs=Presentation(p);assert len(prs.slides)==10
                assert all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
                assert sum(len(s.shapes) for s in prs.slides)>100
            else:assert len(Document(p).paragraphs)>5
            office.append(str(p.relative_to(ROOT)))
    pdfs={}
    for p in ROOT.rglob('*.pdf'):
        d=fitz.open(p);assert len(d)>0
        pdfs[str(p.relative_to(ROOT))]=len(d)
        assert all(page.get_text().strip() for page in d)
        if p.name=='Presentacion.pdf':
            assert len(d)==10
            for pg in d:
                for block in pg.get_text('dict')['blocks']:
                    if block['type']!=0:continue
                    x0,y0,x1,y1=block['bbox']
                    assert x0>=0 and y0>=0 and x1<=pg.rect.width+1 and y1<=pg.rect.height+1,block
        if p.name=='Tarea.pdf':assert len(d)==4
    with (ROOT/'Matriz_temas.csv').open(encoding='utf-8-sig') as f:
        assert len(list(csv.DictReader(f)))==27
    media={}
    for name in ('Video.mp4','Audio.mp3'):
        p=TOPIC/name
        info=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_format','-show_streams','-of','json',str(p)],text=True))
        media[name]={'duration_seconds':float(info['format']['duration']),'bytes':p.stat().st_size,'streams':[{k:s[k] for k in ('codec_name','codec_type','width','height','sample_rate','channels','pix_fmt','r_frame_rate') if k in s} for s in info['streams']]}
    assert 120<=media['Video.mp4']['duration_seconds']<=240
    loudness=json.loads((BUILD/'sonoridad.json').read_text())
    for r in loudness.values():assert float(r['input_tp'])<0
    report={'fecha':'2026-10-29','alcance':'1 piloto: Negocios Electrónicos, 1.1. Las páginas de inicio; otros 26 pendientes de asignación','subtemas_matriz':27,'office_integridad_y_reapertura':office,'pdf_abiertos_paginas':pdfs,'medios':media,'decodificacion_ffmpeg_completa':{'Video.mp4':'sin errores','Audio.mp3':'sin errores'},'reproduccion_chromium':{'Video.mp4':{'duracion':218.125,'avance':0.831067,'error':None,'resolucion':[1920,1080]},'Audio.mp3':{'duracion':197.151995,'avance':0.854677,'error':None}},'cuestionario':{'preguntas':10,'aciertos':'10/10','errores':'0/10','reinicio':True,'teclado':True,'viewport_movil':[390,844],'desbordamiento_horizontal':False,'errores_javascript':[]},'sonoridad_ffmpeg':loudness,'limites':['Sin evaluación auditiva humana de naturalidad o pronunciación','Educaplay preparado sin publicar','26 subtemas no desarrollados','ZIP del piloto, no de todos los temas']}
    (ROOT/'Validacion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'README.md').write_text(README,encoding='utf-8')
    vr=['# Informe de validación · piloto NE 1.1','','Comprobaciones ejecutadas el 29 de octubre de 2026. No se atribuye validación a contenidos que no se han generado.','',
    '| Elemento | Resultado comprobado |','|---|---|',
    '| PPTX | Integridad ZIP/XML sin errores; reabierto con python-pptx y abierto/exportado por LibreOffice; 10 diapositivas editables y notas de narración. |',
    '| DOCX tarea | Reabierto con python-docx y abierto/exportado por LibreOffice; 4 páginas, 6 ejercicios y rúbrica propuesta que suma 20 puntos. |',
    '| Matriz | 27 filas de subtemas en CSV, 2 páginas en Word exportado a PDF. Títulos revisados contra la página 1 de cada sílabo. |',
    '| Video | H.264, yuv420p, 1920×1080, 24 fps y pista AAC 48 kHz; 218,125 segundos. Decodificado de principio a fin sin errores con FFmpeg y reproducido en Chromium con avance temporal. |',
    '| Audio | MP3, 44,1 kHz mono; aproximadamente 197,2 segundos. Decodificado íntegramente sin errores con FFmpeg y reproducido en Chromium con avance temporal. |',
    '| Práctica interactiva | 10 preguntas verificadas en Chromium. Todas correctas: 10/10; todas incorrectas: 0/10. Reinicio, retroalimentación y bloqueo de doble respuesta correctos. Uso de teclado y vista móvil 390×844 sin desbordamiento horizontal. Sin errores JavaScript. |',
    '| Revisión visual | Diapositivas abiertas realmente con LibreOffice y exportadas a PDF; hoja de contacto revisada. Tarea revisada en PDF; se corrigió la paginación y los anchos de la rúbrica. Fotogramas extraídos del MP4 comprobado. |','',
    '## Sonoridad de los archivos finales','',
    '| Archivo | Sonoridad integrada LUFS | Pico verdadero dBTP |','|---|---:|---:|']
    for n,r in loudness.items():vr.append(f'| {n} | {r["input_i"]} | {r["input_tp"]} |')
    vr+=['','Medidas obtenidas con el analizador loudnorm de FFmpeg sobre los archivos finales. Pico verdadero inferior a 0 dBTP en ambos. La música original se generó a nivel reducido y se atenuó con sidechain durante la narración.','',
    '## Límites de la comprobación','',
    '- Reproducción y análisis técnicos no equivalen a escucha humana. La voz neuronal usa ef_dora y fonetización latinoamericana es-419; no se certifica naturalidad, pronunciación o preferencia subjetiva sin revisión auditiva humana.',
    '- Los subtítulos SRT y los incrustados usan segmentación proporcional por frase; no se realizó alineación fonética forzada.',
    '- Educaplay no se publicó ni se probó dentro de una cuenta. Se validó una práctica HTML independiente que usa el mismo banco y se dejaron instrucciones para publicar y revisar la versión del servicio.',
    '- Las pruebas no abarcan todos los reproductores, todas las versiones de Office ni un cumplimiento completo de WCAG.',
    '- Este informe corresponde al piloto entregado, no a los otros 26 subtemas ni a futuras regeneraciones.',
    '- La integridad final del ZIP se verifica por CRC al empaquetar y los archivos internos están inventariados en MANIFIESTO_SHA256.txt.','',
    'Detalles de códecs, duraciones y resultados: Validacion.json. Scripts de pruebas reales: Produccion/validar_interactivo.py y Produccion/validar_multimedia.py.']
    (ROOT/'Validacion.md').write_text('\n'.join(vr),encoding='utf-8')
    prod=ROOT/'Produccion';prod.mkdir(exist_ok=True)
    for name in ('contenido.py','generar_documentos.py','generar_audio.py','montar_media.py','validar_interactivo.py','validar_multimedia.py','empaquetar.py'):
        shutil.copy2(BUILD/name,prod/name)
    (prod/'README.md').write_text(PRODUCTION,encoding='utf-8')
    deps={'python-pptx':'1.0.2','python-docx':'1.2.0','Pillow':'12.3.0','numpy':'2.3.5','PyMuPDF':'1.26.6','pypdf':'6.10.0','kokoro-onnx':'0.4.9','soundfile':'0.13.1','onnxruntime':'1.31.1','playwright':'1.62.0'}
    (prod/'requirements.txt').write_text('\n'.join(f'{k}=={v}' for k,v in deps.items())+'\n')
    hashes=[]
    for p in sorted(ROOT.rglob('*')):
        if p.is_file() and p.name!='MANIFIESTO_SHA256.txt':
            hashes.append(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+str(p.relative_to(ROOT)))
    (ROOT/'MANIFIESTO_SHA256.txt').write_text('\n'.join(hashes)+'\n',encoding='utf-8')
    archive=ROOT.parent/'Proyecto_Educativo_Piloto.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(ROOT.rglob('*')):
            if p.is_file():z.write(p,str(Path(ROOT.name)/p.relative_to(ROOT)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        assert len(z.namelist())==len(hashes)+1
        required=['Presentacion.pptx','Video.mp4','Audio.mp3','Tarea.docx','Actividad_Educaplay.txt']
        assert all(any(n.endswith('/'+r) for n in z.namelist()) for r in required)
    print(json.dumps({'zip':str(archive),'MiB':round(archive.stat().st_size/1024/1024,2),'archivos':len(hashes)+1,'crc':'correcto','subtemas_matriz':27,'pilotos':1},ensure_ascii=False),flush=True)

if __name__=='__main__':main()
