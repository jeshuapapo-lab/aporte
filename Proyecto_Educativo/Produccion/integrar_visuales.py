import sys,time,json,shutil,os
from pathlib import Path
sys.path.insert(0,'/workspace/edu-build')
from generar_documentos import Canvas,INK,TEAL,NAVY,WHITE,PALE,MUTED
from pptx import Presentation
import modelos_bpmn as models
from producir import folder
BASE=Path('/workspace/edu-full')

def chart(t,dest,work):
 p=dest/'Presentacion.pptx';prs=Presentation(p);s=prs.slides[6];note=s.notes_slide.notes_text_frame.text
 for shape in list(s.shapes):s.shapes._spTree.remove(shape._element)
 c=Canvas(s);c.rect(0,0,1600,12,TEAL);c.text(80,48,1400,'ESTUDIO DE CASO · DATOS SIMULADOS · PND 4.4',20,TEAL,True)
 c.text(80,122,1430,'Comparar resultados y límites',53,INK,True)
 metrics=[('Ciclo promedio',[40,24],50,'horas','40 % de reducción'),('Casos con corrección',[12,5],15,'por ciento','7 puntos porcentuales menos'),('Dentro de 48 h',[80,95],100,'por ciento','+15 puntos porcentuales')]
 for j,(label,values,limit,unit,annotation) in enumerate(metrics):
  x=80+j*490;c.rect(x,265,460,488,WHITE,r=16);c.text(x+25,287,410,label,28,INK,True,max_h=80)
  for k,v in enumerate(values):
   bx=x+65+k*175;h=250*v/limit;c.rect(bx,658-h,105,h,NAVY if k==0 else TEAL,r=4);c.text(bx-3,605-h,145,str(v)+(' h' if unit=='horas' else ' %'),36,INK,True)
   c.text(bx-20,676,155,'Inicial' if k==0 else 'Futuro',23,MUTED,True)
  c.text(x+20,715,420,annotation,19,TEAL,True)
 c.rect(80,784,1430,40,PALE,r=8);c.text(98,792,1390,'100 casos por período. Cinco casos futuros llegan a 100 horas: la mejora de la media no elimina excepciones.',19,INK,max_h=27)
 c.text(80,846,1430,'El cálculo sobre datos sintéticos no demuestra un efecto causal en una empresa real.',17,MUTED)
 s.notes_slide.notes_text_frame.text=note;c.im.save(work/'06.png')
 temp=work/'Presentacion_mejorada.pptx';prs.save(temp);os.replace(temp,p)

def bpmn(t,dest,work):
 p=dest/'Presentacion.pptx';prs=Presentation(p);old=prs.slides[4];note=old.notes_slide.notes_text_frame.text
 oldid=prs.slides._sldIdLst[4]
 models.directory_global=work
 model=models.create(True,work);models.render(prs,model,True)
 prs.part.drop_rel(oldid.rId);prs.slides._sldIdLst.remove(oldid)
 new_id=prs.slides._sldIdLst[-1];prs.slides._sldIdLst.remove(new_id);prs.slides._sldIdLst.insert(4,new_id)
 prs.slides[4].notes_slide.notes_text_frame.text=note
 shutil.copy2(work/'Estado_futuro.png',work/'04.png')
 assert len(prs.slides)==10
 temp=work/'Presentacion_mejorada.pptx';prs.save(temp);os.replace(temp,p)

if __name__=='__main__':
 jobs={'PND_3_2':bpmn,'PND_4_4':chart};done=[]
 while jobs:
  for key,fn in list(jobs.items()):
   t=json.loads((BASE/'content'/(key+'.json')).read_text());dest=folder(t);work=BASE/'work'/key
   if (work/'VISUAL_INTEGRADO').exists():done.append(key);del jobs[key];continue
   p=dest/'Presentacion.pptx'
   if not p.exists() or time.time()-p.stat().st_mtime<3:continue
   if (dest/'VALIDADO_PRODUCCION.json').exists():raise RuntimeError('El tema terminó antes de integrar su visual: '+key)
   fn(t,dest,work);(work/'VISUAL_INTEGRADO').write_text(key+'\n');done.append(key);del jobs[key];print('VISUAL INTEGRADO '+key,flush=True)
  time.sleep(1)
 (BASE/'visuales_integrados.json').write_text(json.dumps(done))
