import os,sys,json,math,re,time,hashlib,subprocess,concurrent.futures,functools,unicodedata,textwrap,csv,traceback
from pathlib import Path
import numpy as np
import soundfile as sf
import onnxruntime as ort
from kokoro_onnx import Kokoro
sys.path.insert(0,'/workspace/edu-build')
from generar_documentos import Canvas,setup_doc,table,NAVY,TEAL,CREAM,INK,MUTED,WHITE,PALE,GOLD,CORAL
from pptx import Presentation
from pptx.util import Inches
import montar_media as media

ROOT=Path('/workspace/Proyecto_Educativo_Completo');BASE=Path('/workspace/edu-full');SR=24000;FPS=24
ENGINE=None
def init():
    global ENGINE
    opts=ort.SessionOptions();opts.intra_op_num_threads=2;opts.inter_op_num_threads=1
    session=ort.InferenceSession('/workspace/edu-runtime/models/kokoro-v1.0.onnx',sess_options=opts,providers=['CPUExecutionProvider'])
    class Typed:
        def __getattr__(self,k):return getattr(session,k)
        def run(self,o,i):
            i['speed']=np.array([1.0],dtype=np.float32)
            return session.run(o,i)
    ENGINE=Kokoro.from_session(Typed(),'/workspace/edu-runtime/models/voices-v1.0.bin')

def run(cmd):
    r=subprocess.run(cmd,capture_output=True,text=True)
    if r.returncode:raise RuntimeError(r.stderr[-3500:])
    return r

def slug(s):return re.sub(r'[^A-Za-z0-9]+','_',unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode()).strip('_')
def folder(t):return ROOT/('Negocios_Electronicos' if t['code']=='NE' else 'Procesos_para_Negocios_Digitales')/('Tema_'+t['num'].replace('.','_')+'_'+slug(t['title']))
def tts(text,p):
    digest=hashlib.sha256(text.encode()).hexdigest()
    if p.exists() and p.with_suffix('.json').exists():
        d=json.loads(p.with_suffix('.json').read_text())
        if d.get('hash')==digest:return sf.read(p,dtype='float32')[0],d['caps']
    a=[np.zeros(6000,dtype=np.float32)];caps=[];cursor=.25
    for sentence in re.split(r'(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÑ¿¡])',text):
        if not sentence.strip():continue
        samples,sr=ENGINE.create(sentence,voice='ef_dora',lang='es-419',speed=1.)
        assert sr==SR and np.isfinite(samples).all()
        peak=np.abs(samples).max()
        if peak>.95:samples*=.95/peak
        dur=len(samples)/SR;caps.append(dict(start=cursor,end=cursor+dur,text=sentence));cursor+=dur+.22
        a.extend([samples,np.zeros(round(.22*SR),dtype=np.float32)])
    full=np.concatenate(a+[np.zeros(6000,dtype=np.float32)])
    full=np.pad(full,(0,math.ceil(len(full)/1000)*1000-len(full)))
    sf.write(p,full,SR,subtype='PCM_16');p.with_suffix('.json').write_text(json.dumps(dict(hash=digest,caps=caps),ensure_ascii=False))
    return full,caps

def slides_and_scripts(t,dest,work):
    prs=Presentation();prs.slide_width=Inches(1600/120);prs.slide_height=Inches(900/120)
    short=f"{t['materia']} · U{t['num'][0]} · {t['num']}"
    intro=f"Hoy estudiaremos {t['title'].lower()}. El propósito es {t['objective'][0].lower()+t['objective'][1:]} Este contenido forma parte de {t['materia'].lower()}. Trabajaremos con conceptos y un caso ficticio para tomar decisiones justificadas."
    narr=[intro,'Antes de proponer una solución, considera esta pregunta: '+t['problem']+' Al terminar deberás explicar tu decisión, indicar cómo la comprobarías y reconocer los límites de la evidencia utilizada.',t['sections'][0][1],t['sections'][1][1],t['flowtext'],t['case'],t['sections'][2][1],t['sections'][3][1],
    'Para aplicar lo aprendido, desarrolla la tarea editable. '+t['tasks'][0]+' Después, '+t['tasks'][2][0].lower()+t['tasks'][2][1:]+' Justifica tus decisiones con los conceptos de esta explicación y comprueba tus respuestas en la actividad interactiva.',
    'Conserva estas ideas. '+' '.join(t['summary'])+' El caso y sus datos son ficticios. Revisa las fuentes indicadas en los materiales y continúa con la práctica para transferir estos criterios a otra situación.']
    section_for={2:0,3:1,6:2,7:3}
    for i in range(10):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);c=Canvas(slide);dark=i in [0,9]
        if dark:c.rect(0,0,1600,900,NAVY)
        c.rect(0,0,1600,12,TEAL)
        c.text(80,48,1420,short.upper(),20,GOLD if dark else TEAL,True,max_h=58)
        if i==0:
            size=64 if len(t['title'])<75 else 53
            h=c.text(80,158,1410,t['title'],size,WHITE,True,max_h=275)
            c.text(85,185+h,1310,t['objective'],32,'#D6E4E8',max_h=177)
            for j,label in enumerate(['Comprender','Aplicar','Comprobar']):
                c.rect(85+j*490,654,455,104,'#1D4052',r=16);c.text(112+j*490,681,400,label,30,WHITE,True)
        else:
            title={1:'Una decisión que podrás justificar',4:'Del concepto a la acción',5:'Caso aplicado · datos ficticios',8:'Tu actividad de transferencia',9:'Síntesis y referencias'}.get(i,t['sections'][section_for[i]][0] if i in section_for else '')
            c.text(80,125,1430,title,52,WHITE if dark else INK,True,max_h=146)
            if i==1:
                c.rect(80,303,1430,176,WHITE,r=18);c.text(110,326,1370,t['objective'],34,INK,True,max_h=141)
                c.text(90,533,1400,'Pregunta de entrada',26,TEAL,True)
                c.text(90,589,1400,t['problem'],34,INK,max_h=145)
            elif i in section_for:
                text=t['sections'][section_for[i]][1]
                c.rect(80,300,1428,445,WHITE,r=18);c.rect(80,300,11,445,TEAL)
                c.text(118,331,1340,text,31,INK,max_h=382)
            elif i==4:
                for j,label in enumerate(t['flow']):
                    x=80+j*360;c.rect(x,320,330,156,TEAL,r=15)
                    c.text(x+23,342,285,f'{j+1:02d}',21,WHITE,True)
                    c.text(x+23,387,285,label,27,WHITE,True,max_h=82)
                    if j<3:c.text(x+334,369,30,'›',30,INK,True)
                c.rect(80,520,1428,250,WHITE,r=18);c.text(110,548,1368,t['flowtext'],28,MUTED,max_h=210)
            elif i==5:
                c.rect(80,298,1428,302,WHITE,r=18);c.text(112,326,1364,t['case'],29,INK,max_h=248)
                c.rect(80,628,1428,121,PALE,r=15);c.text(109,649,1370,t['problem'],27,INK,True,max_h=89)
            elif i==8:
                for j,task in enumerate(t['tasks']):
                    y=295+j*112;c.ellipse(80,y+7,50,50,TEAL);c.text(95,y+14,35,str(j+1),24,WHITE,True)
                    c.text(160,y+3,1340,task,27,INK,max_h=98)
            else:
                for j,point in enumerate(t['summary']):
                    c.rect(80,285+j*104,1430,91,'#1D4052',r=12);c.text(110,303+j*104,1370,point,30,WHITE,True,max_h=76)
                source='Base curricular: sílabo UPS adjunto, página 1. Referencia de ampliación: '+t['source']
                c.text(90,627,1410,source,23,'#D6E4E8',max_h=160)
        c.text(80,846,1390,'Material didáctico de elaboración original · Caso y datos ficticios',17,'#C5D6DC' if dark else MUTED)
        c.text(1450,846,100,f'{i+1:02d} / 10',17,'#C5D6DC' if dark else MUTED)
        slide.notes_slide.notes_text_frame.text=narr[i]+'\n\nFuente curricular: '+t['materia']+', '+t['num']+', UPS 2026–2027. Referencia: '+t['source']
        c.im.save(work/f'{i:02d}.png')
    prs.core_properties.title=t['title'];prs.core_properties.subject=t['unidad'];prs.core_properties.author='Proyecto educativo · elaboración asistida por IA'
    prs.save(dest/'Presentacion.pptx')
    return narr

def docs_and_quiz(t,dest):
    d=setup_doc('Tarea · '+t['title']);d.sections[0].header.paragraphs[0].text=t['materia'].upper()+' / '+t['num']
    d.add_heading(t['title'],0);d.add_paragraph('ACTIVIDAD ACADÉMICA · CASO FICTICIO',style='Subtitle')
    d.add_paragraph('Nombre: _____________________  Curso: __________  Fecha: __________')
    d.add_heading('Objetivo',1);d.add_paragraph(t['objective'])
    d.add_heading('Instrucciones',1);d.add_paragraph('Resuelve los cuatro ejercicios en este Word editable. Justifica cada decisión y muestra procedimientos cuando haya cálculos. Puedes incorporar diagramas propios. Tiempo sugerido: 60–90 minutos; no es un requisito fijado por el profesor. Declara fuentes y uso de IA. No uses datos personales reales.')
    d.add_heading('Situación de trabajo',1);d.add_paragraph(t['case'])
    d.add_heading('1. Analizar — 5 puntos',1);d.add_paragraph(t['tasks'][0]);d.add_paragraph('Respuesta argumentada: __________________________________________\n________________________________________________________________\n________________________________________________________________')
    d.add_heading('2. Diseñar — 5 puntos',1);d.add_paragraph(t['tasks'][1]);d.add_paragraph('Propuesta y justificación: _________________________________________\n________________________________________________________________\n________________________________________________________________')
    d.add_page_break();d.add_heading('Aplicar y comprobar',0)
    d.add_heading('3. Aplicar — 5 puntos',1);d.add_paragraph(t['tasks'][2]);d.add_paragraph('Procedimiento, modelo o cálculo: __________________________________\n________________________________________________________________\n________________________________________________________________\n________________________________________________________________\n________________________________________________________________')
    d.add_heading('4. Evaluar — 5 puntos',1);d.add_paragraph(t['tasks'][3])
    table(d,['Decisión / medida','Evidencia o procedimiento','Límite o riesgo'],[['________________','________________','________________'],['________________','________________','________________'],['________________','________________','________________']])
    d.add_heading('Reflexión y autoría',1);d.add_paragraph('¿Qué cambiarías si variara una condición importante del caso?\n________________________________________________________________\nFuentes consultadas: ______________________________________________\nHerramientas e IA utilizadas y verificación realizada: ___________________')
    d.add_page_break();d.add_heading('Rúbrica propuesta — 20 puntos',0)
    d.add_paragraph('Esta escala didáctica no reemplaza las ponderaciones ni una rúbrica oficial del profesor.')
    table(d,['Criterio','Puntos','Evidencia de logro completo'],[
      ['Análisis','5','Identifica el problema específico y usa correctamente los conceptos del tema.'],
      ['Diseño','5','Entrega la propuesta solicitada con relaciones claras y decisiones justificadas.'],
      ['Aplicación','5','Desarrolla el procedimiento, modelo o cálculo pedido de forma trazable.'],
      ['Evaluación','5','Propone comprobaciones observables, reconoce riesgos y límites de la evidencia.']],widths=[1.3,.6,5.1])
    d.add_paragraph('Por criterio: 5 = completo y bien justificado; 4 = correcto con una omisión menor; 3 = parcial con errores u omisiones relevantes; 1–2 = evidencia mínima o poco conectada; 0 = ausente o ajeno al objetivo.')
    d.add_heading('Lista de revisión',1)
    for s in ['Respondí los cuatro ejercicios con evidencias propias.','Diferencié hechos del caso y supuestos que yo agregué.','Mis decisiones pueden comprobarse mediante un procedimiento.','Reconocí límites, riesgos y condiciones que cambiarían el resultado.','Conservé el archivo editable y declaré fuentes y ayuda de IA.']:d.add_paragraph('☐ '+s)
    d.add_heading('Fuentes',1);d.add_paragraph('UPS, sílabo 2026–2027, '+t['materia']+', unidad '+t['num'][0]+', tema '+t['num']+'.\n'+t['source'])
    d.save(dest/'Tarea.docx')
    guide='# Guía docente · '+t['title']+'\n\nObjetivo: '+t['objective']+'\n\n## Orientación del caso\n\n'+t['solution']+'\n\n## Conceptos y evidencias\n\n'+'\n\n'.join('### '+h+'\n'+p for h,p in t['sections'])+'\n\n## Corrección\n\nValorar los cuatro ejercicios con la rúbrica de Tarea.docx. Admitir soluciones alternativas justificadas y coherentes con las restricciones. No asignar puntaje completo a una lista sin explicación. El caso y los datos son ficticios; el criterio del profesor prevalece sobre esta propuesta. Las respuestas y la retroalimentación del cuestionario están en Actividad_Educaplay.txt.\n'
    (dest/'Recursos'/'Guia_docente.md').write_text(guide,encoding='utf-8')
    quiz=[]
    for i,q in enumerate(t['questions']):
        options=list(q['distractors']);answer=i%4;options.insert(answer,q['correct']);quiz.append(dict(q=q['q'],options=options,answer=answer,feedback=q['feedback']))
    text=['ACTIVIDAD PREPARADA PARA EDUCAPLAY — SIN PUBLICAR',t['materia']+' · '+t['num']+' · '+t['title'],'Objetivo: '+t['objective'],'Tipo propuesto: cuestionario de opción múltiple. Ocho preguntas, una respuesta correcta, un punto por acierto, sin descuento ni límite de tiempo. Referencia de práctica: seis de ocho. Estas son propuestas didácticas.','Instrucciones: elige una respuesta y revisa la explicación. El caso y los datos son ficticios.','']
    for i,q in enumerate(quiz):
        text+=[f'{i+1}. '+q['q'],*['   '+'ABCD'[j]+') '+a for j,a in enumerate(q['options'])],'Correcta: '+'ABCD'[q['answer']],'Retroalimentación: '+q['feedback'],'']
    text+=['PUBLICACIÓN CON TU CUENTA AUTORIZADA','1. Accede al sitio oficial https://www.educaplay.com/ e inicia sesión.','2. Crea un cuestionario de opciones múltiples y copia título, objetivo e instrucciones. La interfaz y funciones pueden variar con la cuenta.','3. Incorpora las ocho preguntas, sus opciones y su respuesta correcta. Añade retroalimentación si el formato lo permite.','4. Configura español y las opciones de evaluación disponibles. El CSV adjunto es auxiliar, no un formato de importación verificado del servicio.','5. Recorre una vista de prueba completa, revisando aciertos, errores y puntuación.','6. Publica con la privacidad elegida y copia el enlace real de compartir. Verifica su acceso con permisos de estudiante.','Enlace de actividad publicada: no disponible; no se accedió a una cuenta de Educaplay. No se inventa un enlace.','Práctica local funcional: Actividad_Interactiva.html. Es independiente de Educaplay.']
    (dest/'Actividad_Educaplay.txt').write_text('\n'.join(text),encoding='utf-8')
    with (dest/'Recursos'/'Banco_preguntas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['Pregunta','A','B','C','D','Correcta','Retroalimentacion'])
        for q in quiz:w.writerow([q['q'],*q['options'],'ABCD'[q['answer']],q['feedback']])
    template=Path('/workspace/Proyecto_Educativo/Negocios_Electronicos/Tema_1_1_Las_paginas_de_inicio/Actividad_Interactiva.html').read_text()
    template=re.sub(r'const preguntas=.*?;let indice=',lambda m:'const preguntas='+json.dumps(quiz,ensure_ascii=False)+';let indice=',template,flags=re.S)
    template=template.replace('Las páginas de inicio',t['title']).replace('NEGOCIOS ELECTRÓNICOS · TEMA 1.1',t['materia'].upper()+' · TEMA '+t['num']).replace('max="10"','max="8"').replace('puntos>=8','puntos>=6')
    (dest/'Actividad_Interactiva.html').write_text(template,encoding='utf-8')

def make_media(t,dest,work,narr):
    audio_dir=work/'audio';audio_dir.mkdir(exist_ok=True)
    wave=[];captions=[];cursor=0
    section_audio={}
    for i,text in enumerate(narr):
        a,caps=tts(text,audio_dir/f'{i:02d}.wav');wave.append(a)
        for cap in caps:
            words=cap['text'].split();chunks=[words[j:j+11] for j in range(0,len(words),11)]
            for j,chunk in enumerate(chunks):
                start=cursor+cap['start']+(cap['end']-cap['start'])*j/len(chunks)
                end=cursor+cap['start']+(cap['end']-cap['start'])*(j+1)/len(chunks)
                captions.append(dict(start=start,end=end,text=' '.join(chunk)))
        cursor+=len(a)/SR
    if not 120<=cursor<=240:raise RuntimeError(f'Duración fuera de 2–4 minutos: {cursor} {t["id"]}')
    sf.write(work/'voz.wav',np.concatenate(wave),SR,subtype='PCM_16')
    intro='Esta es una explicación para escuchar de manera independiente sobre '+t['title'].lower()+'. No necesitas consultar imágenes. '+t['objective']+' Piensa en cómo resolverías una situación similar en una organización que conozcas; distinguiremos conceptos, decisiones y comprobaciones.'
    outro='Antes de terminar, responde mentalmente: '+t['problem']+' Una orientación posible es la siguiente. '+t['solution']+' Recuerda: '+' '.join(t['summary'])
    ai,_=tts(intro,audio_dir/'intro_audio.wav');ao,_=tts(outro,audio_dir/'cierre_audio.wav')
    podcast=np.concatenate([ai,*[np.concatenate([wave[i],np.zeros(12000,dtype=np.float32)]) for i in [2,3,4,5,6,7]],ao])
    sf.write(work/'podcast.wav',podcast,SR,subtype='PCM_16')
    (dest/'Recursos'/'Guion_video.md').write_text('# Guion del video\n\n'+'\n\n'.join(f'## Escena {i+1}\n\n{s}' for i,s in enumerate(narr)),encoding='utf-8')
    (dest/'Recursos'/'Guion_audio.md').write_text('# Explicación independiente\n\n'+intro+'\n\n'+'\n\n'.join(narr[i] for i in [2,3,4,5,6,7])+'\n\n'+outro,encoding='utf-8')
    (dest/'Subtitulos.srt').write_text('\n\n'.join(f'{i+1}\n{media.srt_time(c["start"])} --> {media.srt_time(c["end"])}\n{c["text"]}' for i,c in enumerate(captions))+'\n',encoding='utf-8')
    (work/'subtitulos.ass').write_text(media.ass(captions),encoding='utf-8')
    media.A=work;media.music(cursor)
    run(['ffmpeg','-v','error','-y','-i',str(work/'voz.wav'),'-af','loudnorm=I=-16:TP=-2:LRA=7','-ar','48000','-c:a','pcm_s24le',str(work/'voz_normalizada.wav')])
    run(['ffmpeg','-v','error','-y','-i',str(work/'podcast.wav'),'-af','loudnorm=I=-16:TP=-1.5:LRA=7','-ar','44100','-c:a','libmp3lame','-q:a','2','-metadata','title='+t['title']+' — explicación independiente',str(dest/'Audio.mp3')])
    lines=[]
    for i,a in enumerate(wave):lines.extend([f"file '{work/f'{i:02d}.png'}'",f'duration {len(a)/SR:.8f}'])
    lines.append(f"file '{work/'09.png'}'");(work/'imagenes.txt').write_text('\n'.join(lines)+'\n')
    vf=f"[0:v]fps=24,pad=1920:1080:160:0:color=0x112B3C[base];color=c=0x087E8B:s=1600x6:r=24:d={cursor}[bar];[base][bar]overlay=x='160-1600+1600*min(t/{cursor},1)':y=914:shortest=1,drawbox=x=0:y=914:w=160:h=6:color=0x112B3C:t=fill,subtitles={work/'subtitulos.ass'},fade=t=in:st=0:d=0.3,fade=t=out:st={cursor-.3}:d=0.3,format=yuv420p[v];[1:a]asplit=2[voz][control];[2:a][control]sidechaincompress=threshold=0.05:ratio=5:attack=20:release=400[m];[voz][m]amix=inputs=2:duration=first:normalize=0,alimiter=limit=0.94:level=0[a]"
    run(['ffmpeg','-v','error','-y','-f','concat','-safe','0','-i',str(work/'imagenes.txt'),'-i',str(work/'voz_normalizada.wav'),'-i',str(work/'musica_original.wav'),'-filter_complex_threads','1','-filter_complex',vf,'-map','[v]','-map','[a]','-t',str(cursor),'-c:v','libx264','-preset','veryfast','-crf','22','-threads','1','-r','24','-c:a','aac','-b:a','160k','-movflags','+faststart','-metadata','title='+t['title'],str(dest/'Video.mp4')])
    return cursor,len(podcast)/SR

def produce(path):
    start=time.time();t=json.loads(Path(path).read_text());dest=folder(t);work=BASE/'work'/t['id'];work.mkdir(parents=True,exist_ok=True);dest.mkdir(parents=True,exist_ok=True);(dest/'Recursos').mkdir(exist_ok=True)
    if (dest/'VALIDADO_PRODUCCION.json').exists():return {'id':t['id'],'cached':True}
    narr=slides_and_scripts(t,dest,work);docs_and_quiz(t,dest)
    print(t['id']+' documentos creados; sintetizando narraciones',flush=True)
    video,audio=make_media(t,dest,work,narr)
    result=dict(id=t['id'],title=t['title'],folder=str(dest),video_seconds=video,audio_seconds=audio,seconds_elapsed=time.time()-start)
    (dest/'VALIDADO_PRODUCCION.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print('TERMINADO '+json.dumps(result,ensure_ascii=False),flush=True);return result

if __name__=='__main__':
    futures={};seen=set();failures=[]
    with concurrent.futures.ProcessPoolExecutor(max_workers=2,initializer=init) as pool:
        while True:
            for p in sorted((BASE/'content').glob('*.json')):
                if str(p) not in seen:
                    seen.add(str(p));futures[pool.submit(produce,str(p))]=str(p)
            for future in list(futures):
                if future.done():
                    path=futures.pop(future)
                    try:future.result()
                    except Exception as e:
                        failures.append(dict(path=path,error=str(e)));print('ERROR '+path+' '+str(e),flush=True)
            (BASE/'estado.json').write_text(json.dumps(dict(submitted=len(seen),pending=len(futures),failed=failures),ensure_ascii=False,indent=2))
            if (BASE/'CATALOGO_COMPLETO').exists() and not futures:break
            time.sleep(2)
    if failures:raise SystemExit('Hay temas fallidos: revisar estado.json')
