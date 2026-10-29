import sys,json,copy,xml.etree.ElementTree as ET,math
from lxml import etree
from pathlib import Path
sys.path.insert(0,'/workspace/edu-build')
from generar_documentos import Canvas,NAVY,TEAL,CREAM,INK,MUTED,WHITE,PALE,GOLD
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
BASE=Path('/workspace/edu-full');ROOT=Path('/workspace/Proyecto_Educativo_Completo')
NS={'bpmn':'http://www.omg.org/spec/BPMN/20100524/MODEL','bpmndi':'http://www.omg.org/spec/BPMN/20100524/DI','dc':'http://www.omg.org/spec/DD/20100524/DC','di':'http://www.omg.org/spec/DD/20100524/DI','xsi':'http://www.w3.org/2001/XMLSchema-instance'}
for k,v in NS.items():ET.register_namespace(k,v)
def node(p,kind,attr=None,text=None,ns='bpmn'):
 e=ET.SubElement(p,'{'+NS[ns]+'}'+kind,attr or {})
 if text is not None:e.text=text
 return e
def create(future,directory):
 name='Estado_futuro' if future else 'Estado_actual'
 definitions=ET.Element('{'+NS['bpmn']+'}definitions',{'id':'Definitions_'+name,'targetNamespace':'https://example.invalid/educativo/bpmn'})
 process=node(definitions,'process',dict(id='Proceso_'+name,isExecutable='false',name='Soporte de solicitudes · '+name.replace('_',' ').lower()))
 positions={'Inicio':(125,410,36,36),'Registrar':(220,386,165,84),'Verificar':(435,386,180,84),'Decision':(670,401,56,56),'Preparar':(805,386,182,84),'Responder':(1055,386,180,84),'Fin':(1335,410,36,36),'Pedir':(805,600,182,84),'Recibir':(1055,600,180,84)}
 elements=[('Inicio','startEvent','Solicitud recibida'),('Registrar','task','Registrar y asignar' if future else 'Registrar solicitud'),('Verificar','task','Validar al recibir' if future else 'Revisar requisitos'),('Decision','exclusiveGateway','¿Completa?'),('Preparar','task','Preparar respuesta'),('Responder','sendTask','Comunicar respuesta'),('Fin','endEvent','Respuesta útil'),('Pedir','sendTask','Solicitar datos'),('Recibir','receiveTask','Recibir datos')]
 flows=[('F1','Inicio','Registrar',''),('F2','Registrar','Verificar',''),('F3','Verificar','Decision',''),('F4','Decision','Preparar','Sí'),('F5','Preparar','Responder',''),('F6','Responder','Fin',''),('F7','Decision','Pedir','No'),('F8','Pedir','Recibir',''),('F9','Recibir','Verificar','')]
 coords={'F1':[(161,428),(220,428)],'F2':[(385,428),(435,428)],'F3':[(615,428),(670,428)],'F4':[(726,428),(805,428)],'F5':[(987,428),(1055,428)],'F6':[(1235,428),(1335,428)],'F7':[(698,457),(698,642),(805,642)],'F8':[(987,642),(1055,642)],'F9':[(1145,684),(1145,738),(525,738),(525,470)]}
 for eid,typ,label in elements:
  e=node(process,typ,dict(id=eid,name=label))
  if typ=='exclusiveGateway':e.set('default','F7')
  for fid,source,target,_ in flows:
   if target==eid:node(e,'incoming',text=fid)
   if source==eid:node(e,'outgoing',text=fid)
 for fid,source,target,label in flows:
  e=node(process,'sequenceFlow',dict(id=fid,sourceRef=source,targetRef=target))
  if label:e.set('name',label)
  if fid=='F4':node(e,'conditionExpression',{'{'+NS['xsi']+'}type':'bpmn:tFormalExpression'},'solicitudCompleta = true')
 diagram=node(definitions,'BPMNDiagram',dict(id='Diagram_'+name),ns='bpmndi');plane=node(diagram,'BPMNPlane',dict(id='Plane_'+name,bpmnElement='Proceso_'+name),ns='bpmndi')
 for eid,(x,y,w,h) in positions.items():
  shape=node(plane,'BPMNShape',dict(id='Shape_'+eid,bpmnElement=eid),ns='bpmndi');node(shape,'Bounds',dict(x=str(x),y=str(y),width=str(w),height=str(h)),ns='dc')
 for fid,points in coords.items():
  edge=node(plane,'BPMNEdge',dict(id='Edge_'+fid,bpmnElement=fid),ns='bpmndi')
  for x,y in points:node(edge,'waypoint',dict(x=str(x),y=str(y)),ns='di')
 path=directory/(name+'.bpmn');ET.ElementTree(definitions).write(path,encoding='utf-8',xml_declaration=True)
 # Validación estructural: referencias únicas y flujos dentro del mismo proceso.
 tree=ET.parse(path);ids={e.get('id') for e in tree.iter() if e.get('id')};assert len(ids)==len([e for e in tree.iter() if e.get('id')])
 for e in tree.findall('.//bpmn:sequenceFlow',NS):assert e.get('sourceRef') in ids and e.get('targetRef') in ids
 return positions,elements,flows,coords,name
def render(prs,model,future):
 positions,elements,flows,coords,name=model
 s=prs.slides.add_slide(prs.slide_layouts[6]);c=Canvas(s)
 c.text(80,45,1440,'BPMN · MODELO DE ENSEÑANZA · NO EJECUTABLE',20,TEAL,True)
 c.text(80,107,1440,'Propuesta de atención' if future else 'Proceso de referencia',55,INK,True)
 description='Asignación explícita y validación temprana. El resultado debe comprobarse en un piloto.' if future else 'Versión simplificada del recorrido; las esperas y responsabilidades deben investigarse.'
 c.text(80,202,1430,description,27,MUTED,max_h=76)
 c.rect(80,326,1430,456,WHITE,r=12);c.text(104,344,1360,'EQUIPO DE SOPORTE · UN SOLO PARTICIPANTE',18,TEAL,True)
 # Aristas antes de los nodos para preservar etiquetas y formas.
 for fid,points in coords.items():
  c.d.line(points,fill=INK,width=3)
  line=s.shapes.add_connector(1, Inches(points[0][0]/120),Inches(points[0][1]/120),Inches(points[-1][0]/120),Inches(points[-1][1]/120))
  # Las rutas con quiebres se dibujan como segmentos editables.
  s.shapes._spTree.remove(line._element)
  for (x1,y1),(x2,y2) in zip(points,points[1:]):
   ln=s.shapes.add_connector(1,Inches(x1/120),Inches(y1/120),Inches(x2/120),Inches(y2/120));ln.line.color.rgb=RGBColor.from_string(INK[1:])
   if (x2,y2)==points[-1]:ln._element.spPr.get_or_add_ln().append(etree.Element('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd',type='triangle',w='med',len='med'))
  x,y=points[-1];px,py=points[-2];ang=math.atan2(y-py,x-px)
  tip=[(x,y),(x-12*math.cos(ang-.4),y-12*math.sin(ang-.4)),(x-12*math.cos(ang+.4),y-12*math.sin(ang+.4))];c.d.polygon(tip,fill=INK)
 for eid,typ,label in elements:
  x,y,w,h=positions[eid]
  if typ in ['startEvent','endEvent']:
   c.ellipse(x,y,w,h,TEAL)
   if typ=='startEvent':c.ellipse(x+2,y+2,w-4,h-4,WHITE)
   if typ=='endEvent':c.ellipse(x+7,y+7,w-14,h-14,WHITE)
   c.text(x-34,y-49,115,label,16,INK,True,max_h=49)
  elif typ=='exclusiveGateway':
   sh=s.shapes.add_shape(MSO_SHAPE.DIAMOND,Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(PALE[1:]);sh.line.color.rgb=RGBColor.from_string(TEAL[1:]);c.d.polygon([(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)],fill=PALE,outline=TEAL)
   c.text(x+17,y+10,w-20,'X',25,INK,True);c.text(x-25,y-43,130,label,18,INK,True)
  else:
   c.rect(x,y,w,h,PALE,r=13,line=TEAL);c.text(x+12,y+14,w-24,label,19,INK,True,max_h=67)
 c.text(742,398,60,'Sí',18,TEAL,True);c.text(712,535,60,'No',18,TEAL,True)
 c.text(80,820,1435,'Ruta exclusiva: una solicitud incompleta vuelve a revisión cuando se reciben datos. El caso sin respuesta requiere una regla adicional según el alcance.',18,MUTED,max_h=50)
 c.im.save(directory_global/(name+'.png'))
 s.notes_slide.notes_text_frame.text=description+'\nModelo BPMN 2.0.2 educativo y no ejecutable. La rama No es el flujo por defecto. El tratamiento de abandono o falta de respuesta debe acordarse antes de usarlo como especificación operacional completa.'
 return c.im

if __name__=='__main__':
 directory_global=BASE/'modelos';directory_global.mkdir(exist_ok=True)
 prs=Presentation();prs.slide_width=Inches(1600/120);prs.slide_height=Inches(900/120)
 for future in [False,True]:render(prs,create(future,directory_global),future)
 prs.save(directory_global/'Modelos_BPMN_editables.pptx')
 (directory_global/'LEEME.md').write_text('# Modelos BPMN de enseñanza\n\nDos archivos BPMN 2.0.2 no ejecutables: referencia y propuesta futura. Contienen elementos y diagramación XML; pueden importarse en herramientas compatibles. La sintaxis XML y las referencias internas se comprobaron, pero no se ejecutó un motor de procesos ni una validación completa del esquema OMG. Los nombres y las reglas son de un caso ficticio. La rama No es la ruta por defecto y la rama Sí tiene una condición explícita. El modelo simplifica el alcance y no incluye un temporizador de abandono: el equipo debe definir ese tratamiento antes de considerarlo una especificación operacional completa.\n\nLos PNG permiten visualizar. El PPTX contiene formas y textos editables. Las flechas con quiebres son segmentos editables de dibujo; el significado formal de los flujos se conserva en los archivos BPMN.\n',encoding='utf-8')
 print('Modelos BPMN y presentación editables generados.')
