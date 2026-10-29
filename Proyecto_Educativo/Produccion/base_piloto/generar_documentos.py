import csv, html, json, math, shutil, sys, textwrap
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from contenido import *

NAVY='#112B3C'; TEAL='#087E8B'; CORAL='#D9563F'; CREAM='#F5F2EB'; INK='#18394A'; MUTED='#536975'; WHITE='#FFFFFF'; PALE='#E2F0EE'; GOLD='#E7AE53'
FONT='/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf'
BOLD='/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf'

def rgb(color): return RGBColor.from_string(color.lstrip('#'))

class Canvas:
    def __init__(self, slide):
        self.slide=slide
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=rgb(CREAM)
        self.im=Image.new('RGB',(1600,900),CREAM)
        self.d=ImageDraw.Draw(self.im)
    def rect(self,x,y,w,h,color,r=0,line=None):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x/120), Inches(y/120), Inches(w/120), Inches(h/120))
        shape.fill.solid();shape.fill.fore_color.rgb=rgb(color)
        if line: shape.line.color.rgb=rgb(line)
        else: shape.line.fill.background()
        if r:
            try: shape.adjustments[0]=min(.2,r/min(w,h))
            except Exception: pass
            self.d.rounded_rectangle((x,y,x+w,y+h),radius=r,fill=color,outline=line)
        else: self.d.rectangle((x,y,x+w,y+h),fill=color,outline=line)
    def ellipse(self,x,y,w,h,color):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x/120), Inches(y/120), Inches(w/120), Inches(h/120))
        shape.fill.solid();shape.fill.fore_color.rgb=rgb(color);shape.line.fill.background()
        self.d.ellipse((x,y,x+w,y+h),fill=color)
    def text(self,x,y,w,text,size=30,color=INK,bold=False,max_h=None):
        font=ImageFont.truetype(BOLD if bold else FONT,size)
        lines=[]
        for para in text.split('\n'):
            if not para: lines.append('');continue
            line=''
            for word in para.split():
                candidate=(line+' '+word).strip()
                if font.getlength(candidate)>w and line:
                    lines.append(line);line=word
                else: line=candidate
            lines.append(line)
        lh=round(size*1.38)
        height=lh*len(lines)
        if max_h is not None and height>max_h: raise ValueError(f'Text overflow {height}>{max_h}: {text[:80]}')
        if y+height>873: raise ValueError(f'Slide bottom overflow: {text[:80]}')
        for i,line in enumerate(lines):self.d.text((x,y+i*lh),line,font=font,fill=color,anchor='lt')
        box=self.slide.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches((height+14)/120))
        tf=box.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=0;tf.margin_top=0;tf.margin_right=0;tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name='Noto Sans';p.font.size=Pt(size*.6);p.font.bold=bold;p.font.color.rgb=rgb(color)
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(lh*.6)
        return height
    def bean(self,x,y,scale=1):
        self.ellipse(x,y,180*scale,250*scale,CORAL)
        self.ellipse(x+24*scale,y+22*scale,132*scale,205*scale,'#B44231')
        self.rect(x+85*scale,y+15*scale,7*scale,222*scale,GOLD,r=3)
        for j in range(4):self.rect(x+48*scale,y+(55+j*44)*scale,85*scale,5*scale,GOLD,r=2)
        self.ellipse(x+84*scale,y-37*scale,130*scale,45*scale,TEAL)

def card(c,x,y,w,h,title,body,number=None,dark=False):
    c.rect(x,y,w,h,NAVY if dark else WHITE,r=18)
    offset=0
    if number is not None:
        c.text(x+28,y+27,w-56,str(number).zfill(2),23,TEAL if not dark else GOLD,True)
        offset=43
    c.text(x+28,y+29+offset,w-56,title,30,WHITE if dark else INK,True,max_h=94)
    c.text(x+28,y+91+offset,w-56,body,27,'#D6E4E8' if dark else MUTED,max_h=h-114-offset)

def make_slides():
    TOPIC.mkdir(parents=True,exist_ok=True)
    (TOPIC/'Recursos').mkdir(exist_ok=True)
    (BUILD/'slides').mkdir(exist_ok=True)
    prs=Presentation();prs.slide_width=Inches(1600/120);prs.slide_height=Inches(900/120)
    for i,s in enumerate(SLIDES):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);c=Canvas(slide)
        dark=s['kind'] in ('cover','conclusion')
        if dark:c.rect(0,0,1600,900,NAVY)
        c.rect(0,0,1600,12,TEAL)
        c.text(80,50,1400,s['tag'],22,GOLD if dark else TEAL,True)
        if s['kind']=='cover':
            c.text(80,143,1030,s['title'],94,WHITE,True)
            c.text(85,456,850,s['lead'],38,'#D6E4E8')
            c.rect(85,610,870,110,'#1D4052',r=15)
            c.text(109,630,822,s['note'],26,WHITE,max_h=82)
            c.ellipse(1120,240,340,340,'#1D4052');c.bean(1190,285,1.05)
            c.text(1090,660,440,'Una oferta clara.\nUn siguiente paso útil.',29,WHITE,True)
        else:
            c.text(80,106,1450,s['title'],61,WHITE if dark else INK,True,max_h=176)
            kind=s['kind']
            if kind=='cards':
                for j,(a,b) in enumerate(s['cards']):card(c,80+j*490,320,462,345,a,b,j+1)
            elif kind=='anatomy':
                c.rect(80,310,690,438,WHITE,r=18,line='#D1DCDA')
                c.rect(98,328,654,49,NAVY,r=6);c.text(117,339,618,'MINGA CACAO       Inicio   Regalos   Contacto',18,WHITE,True)
                c.rect(114,398,622,167,PALE,r=10)
                c.text(134,417,404,'Cacao ecuatoriano\npara regalar',28,INK,True)
                c.rect(134,505,245,38,TEAL,r=6);c.text(151,512,212,'Explorar regalos',18,WHITE,True)
                c.bean(619,423,.37)
                for j,label in enumerate(['Oferta destacada','Confianza verificable']):
                    c.rect(114+j*316,584,302,81,CREAM,r=8);c.text(132+j*316,605,270,label,20,INK,True)
                c.rect(114,684,622,43,NAVY,r=6);c.text(130,695,585,'Contacto · Ayuda · Información legal',17,WHITE)
                for j,(a,b) in enumerate(s['items']):
                    yy=314+j*114
                    c.ellipse(825,yy+2,43,43,TEAL);c.text(837,yy+7,28,str(j+1),22,WHITE,True)
                    c.text(890,yy,615,a,27,INK,True)
                    c.text(890,yy+43,615,b,24,MUTED,max_h=66)
            elif kind=='comparison':
                card(c,80,315,695,425,*s['left']);card(c,805,315,710,425,*s['right'],dark=True)
                c.rect(80,315,9,425,CORAL);c.rect(805,315,9,425,TEAL)
            elif kind=='journey':
                for j,(a,b) in enumerate(s['steps']):
                    x=80+j*490;c.rect(x,317,462,128,TEAL,r=14);c.text(x+26,335,410,a,31,WHITE,True);c.text(x+26,392,410,b,23,WHITE)
                for j,(a,b) in enumerate(s['cards']):card(c,80+j*735,480,707,228,a,b)
            elif kind=='checklist':
                for j,(a,b) in enumerate(s['items']):card(c,80+(j%2)*735,313+(j//2)*204,707,192,a,b)
            elif kind=='metric':
                c.rect(80,315,780,150,NAVY,r=16);c.text(110,355,725,s['formula'],50,WHITE,True)
                c.text(90,495,756,s['definition'],28,MUTED,max_h=176)
                for j,(a,b) in enumerate(s['cards']):card(c,900,313+j*220,614,205,a,b)
            elif kind=='challenge':
                for j,(a,b) in enumerate(s['items']):card(c,80+j*490,315,462,390,a,b,j+1)
            elif kind=='conclusion':
                for j,(a,b) in enumerate(s['cards']):card(c,80+j*490,320,462,345,a,b,j+1,dark=True)
            elif kind=='sources':
                for j,(a,b) in enumerate(s['items']):
                    yy=310+j*133;c.text(85,yy,400,a,27,TEAL,True);c.text(500,yy,996,b,27,MUTED,max_h=113)
            if s.get('takeaway'):
                c.rect(80,767,1435,52,'#1D4052' if dark else PALE,r=9)
                c.text(98,779,1400,s['takeaway'],18,WHITE if dark else INK,max_h=35)
        c.text(80,850,1400,SUBTITLE+'   /   Material de elaboración original',17,'#C5D6DC' if dark else MUTED)
        c.text(1456,850,93,f'{i+1:02d} / {len(SLIDES):02d}',17,'#C5D6DC' if dark else MUTED)
        slide.notes_slide.notes_text_frame.text='Guion de narración:\n'+s['narration']+'\n\nEl caso Minga Cacao y los datos son ficticios. Base curricular: sílabo UPS adjunto, p. 1. Accesibilidad: W3C, WCAG 2.2 (2023).'
        c.im.save(BUILD/'slides'/f'{i:02d}.png')
    prs.core_properties.title=TITLE;prs.core_properties.subject=SUBTITLE
    prs.core_properties.author='Proyecto educativo · elaboración asistida por IA'
    prs.save(TOPIC/'Presentacion.pptx')
    print('PPTX generado: 10 diapositivas editables y notas',flush=True)

def setup_doc(title):
    d=Document();sec=d.sections[0];sec.top_margin=DI(.65);sec.bottom_margin=DI(.65);sec.left_margin=DI(.75);sec.right_margin=DI(.75)
    style=d.styles['Normal'];style.font.name='Noto Sans';style.font.size=DP(10);style.paragraph_format.space_after=DP(6)
    for name,size in [('Title',30),('Heading 1',18),('Heading 2',12)]:
        d.styles[name].font.name='Noto Sans';d.styles[name].font.size=DP(size);d.styles[name].font.color.rgb=DC.from_string(INK[1:])
        d.styles[name].paragraph_format.space_before=DP(10)
        d.styles[name].paragraph_format.space_after=DP(5)
    h=sec.header.paragraphs[0];h.text='NEGOCIOS ELECTRÓNICOS  /  UNIDAD 1  /  TEMA 1.1';h.style=d.styles['Caption']
    f=sec.footer.paragraphs[0];f.text='Proyecto educativo · caso ficticio · '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    d.core_properties.title=title;return d

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=str(txt)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:run.font.size=DP(9)
    if widths:
        t.autofit=False
        for col,width in zip(t.columns,widths):col.width=DI(width)
        for row in t.rows:
            for cell,width in zip(row.cells,widths):cell.width=DI(width)
    return t

def make_task():
    d=setup_doc('Tarea · Las páginas de inicio')
    d.add_heading('Una entrada que oriente',0)
    d.add_paragraph('TAREA ACADÉMICA · LAS PÁGINAS DE INICIO',style='Subtitle')
    d.add_paragraph('Nombre: _____________________  Curso: __________  Fecha: __________')
    d.add_heading('Objetivo',1)
    d.add_paragraph('Proponer y justificar una página de inicio que comunique la oferta de un negocio electrónico, jerarquice los recorridos y pueda evaluarse con criterios de accesibilidad y resultados definidos.')
    d.add_heading('Instrucciones de entrega',1)
    d.add_paragraph('Trabajo individual: 60–90 minutos estimados, no fijados por el profesor. Completa este Word editable e incluye un boceto propio o una fotografía legible de tu dibujo. No necesitas programar. Justifica las decisiones y declara fuentes y ayuda de IA. No inventes certificaciones ni reseñas.')
    d.add_heading('Caso ficticio: Minga Cacao',1)
    d.add_paragraph('Un negocio de Guayaquil ofrece chocolates de cacao ecuatoriano pensados para regalar. Su inicio solo dice «Bienvenidos», muestra una fotografía sin explicación y contiene seis botones idénticos con la etiqueta «Más». Las personas no reconocen con facilidad qué se vende ni por dónde continuar. No hay información confirmada sobre plazos de entrega, certificaciones o testimonios.')
    d.add_heading('1. Diagnóstico — 3 puntos',1)
    d.add_paragraph('Identifica tres problemas distintos. Para cada uno, explica qué dificultad produce en una persona que busca un regalo y qué cambio propondrías.')
    table(d,['Problema observado','Efecto en el visitante','Cambio propuesto'],[['1. __________________','__________________','__________________'],['2. __________________','__________________','__________________'],['3. __________________','__________________','__________________']])
    d.add_heading('2. Oferta y acción — 4 puntos',1)
    d.add_paragraph('Redacta un título principal de hasta 14 palabras y un texto de apoyo de hasta 35 palabras. Elige una acción prioritaria con una etiqueta de 2–5 palabras y explica a qué destino debe llevar. No prometas condiciones que no figuran en el caso.')
    d.add_paragraph('Título: __________________________________________________________\nApoyo: _________________________________________________________\nBotón y destino: __________________________________________________')
    d.add_page_break()
    d.add_heading('Del mensaje al recorrido',0)
    d.add_heading('3. Boceto móvil y jerarquía — 5 puntos',1)
    d.add_paragraph('Construye un boceto de una página de inicio para una pantalla móvil de 390 píxeles de ancho. Es una referencia de diseño, no una dimensión obligatoria del sílabo. Incluye identidad/navegación, propuesta y acción principal, contenido de apoyo y pie de página. Usa etiquetas comprensibles; marca qué información debe verificarse antes de publicar.')
    table(d,['Zona del boceto','Contenido, orden y justificación'],[
        ['Identidad y navegación','[Completa aquí: nombre, accesos y prioridades]'],
        ['Propuesta y acción principal','[Completa aquí: mensaje, imagen prevista, botón y destino]'],
        ['Contenido de apoyo','[Completa aquí: oferta y señales verificables]'],
        ['Pie de página','[Completa aquí: contacto, ayuda e información legal]']])
    d.add_paragraph('Inserta tu dibujo o imagen de boceto a continuación. Puedes ampliar este espacio al editar.\n\n\n\n\n')
    d.add_heading('4. Recorridos de dos visitantes — 2 puntos',1)
    d.add_paragraph('Describe en 3–4 pasos cómo continuarían: a) alguien que ya sabe qué regalo busca; b) alguien que aún está comparando. Explica cómo tu jerarquía ayuda a ambos sin presentar todas las acciones con igual importancia.')
    d.add_paragraph('Recorrido A: _____________________________________________________\nRecorrido B: _____________________________________________________\nJustificación: ____________________________________________________')
    d.add_page_break()
    d.add_heading('Comprobar el diseño',0)
    d.add_heading('5. Accesibilidad aplicada — 3 puntos',1)
    d.add_paragraph('Propón dos mejoras concretas e indica cómo comprobarlas. Considera teclado/foco, adaptación móvil, contraste, etiquetas o alternativas de imágenes. No afirmes que todo el sitio cumple WCAG solo por superar dos comprobaciones.')
    table(d,['Mejora específica','Procedimiento de comprobación','Resultado esperado'],[['1. __________________','__________________','__________________'],['2. __________________','__________________','__________________']])
    d.add_paragraph('Referencia: WCAG 2.2, criterio 1.4.3: 4,5:1 para texto normal y 3:1 para texto grande, sujeto a sus excepciones. Fuente: https://www.w3.org/TR/WCAG22/')
    d.add_heading('6. Medición e interpretación — 3 puntos',1)
    d.add_paragraph('Datos hipotéticos del mes observado: 2.000 sesiones vieron el inicio; 50 de esas sesiones completaron una solicitud de información, registrada una sola vez por sesión. Para este ejercicio, completar la solicitud es la acción definida.')
    d.add_paragraph('a) Calcula la proporción de sesiones que completaron la acción. Muestra fórmula, sustitución y resultado porcentual.\nb) Explica por qué el resultado no debe describirse como porcentaje de ventas.\nc) Define qué medirías al probar tu rediseño: acción, período y conjunto de sesiones. Señala una precaución para evitar atribuir causalidad sin evidencia.')
    d.add_paragraph('Cálculo: _________________________________________________________\nInterpretación: ___________________________________________________\nPlan de medición y precaución: _____________________________________\n________________________________________________________________')
    d.add_heading('Declaración de autoría y uso de IA',1)
    d.add_paragraph('Herramientas utilizadas: ___________________________________________\nQué sugerencias revisé o modifiqué: __________________________________\nFuentes consultadas: ______________________________________________')
    d.add_page_break()
    d.add_heading('Criterios de evaluación',0)
    d.add_paragraph('Rúbrica propuesta para este material: 20 puntos. Es una escala didáctica local; no reemplaza las ponderaciones oficiales ni una rúbrica que el profesor entregue después.')
    table(d,['Criterio','Máximo','Evidencia para el puntaje completo'],[
        ['Diagnóstico','3','Tres problemas distintos, cada uno con efecto y mejora pertinente.'],
        ['Oferta y acción','4','Mensaje específico y veraz, apoyo coherente, acción descriptiva y destino justificado.'],
        ['Boceto y jerarquía','5','Cuatro zonas comprensibles, orden móvil viable y justificación vinculada al público.'],
        ['Recorridos','2','Dos recorridos coherentes con intenciones diferentes.'],
        ['Accesibilidad','3','Dos mejoras concretas con comprobación y resultado observable.'],
        ['Medición','3','Cálculo correcto, distinción entre solicitud y venta, y plan delimitado con precaución causal.']],widths=[1.5,.6,4.9])
    d.add_paragraph('Para cada criterio: evidencia completa y justificada = 100 % de su puntaje; mayormente correcta con una omisión relevante = 75 %; respuesta parcial con justificación insuficiente = 50 %; respuesta mínima = 25 %; ausencia o respuesta ajena al criterio = 0 %. Se permiten cuartos de punto.')
    d.add_heading('Antes de entregar',1)
    for t in ['El título explica la oferta sin inventar promesas.','El botón principal indica una acción y un destino.','El boceto y los dos recorridos son coherentes.','Las mejoras incluyen procedimientos de comprobación.','El cálculo y su interpretación usan la misma acción y denominador.','El archivo es editable y las fuentes y el uso de IA están declarados.']:
        d.add_paragraph('☐ '+t)
    d.add_heading('Base curricular y referencia técnica',1)
    d.add_paragraph(SOURCES[0][0]+'\n'+SOURCES[2][0]+' '+SOURCES[2][1])
    d.save(TOPIC/'Tarea.docx');print('DOCX editable generado',flush=True)
    (TOPIC/'Recursos'/'Guia_docente.md').write_text('''# Guía docente · tema 1.1

Objetivo: transferir criterios de orientación a una propuesta de página de inicio. La alineación curricular corresponde a «Conoce la estructura de un negocio electrónico», unidad 1. Los objetivos, la tarea y la rúbrica son propuestas didácticas propias, no instrucciones atribuidas al profesor.

## Secuencia sugerida
1. Activar conocimientos: pedir que una persona describa la oferta de una web conocida (3 minutos).
2. Usar la presentación o el video. Pedir una razón detrás de cada elección de mensaje y recorrido (10–15 minutos).
3. Resolver la tarea editable (60–90 minutos estimados).
4. Practicar el cuestionario y revisar retroalimentación (8–12 minutos).

## Orientaciones para corregir
- Diagnóstico: mensaje «Bienvenidos» sin oferta; seis botones «Más» sin jerarquía ni destino claro; fotografía sin explicación. Aceptar problemas alternativos argumentados, sin suponer condiciones no dadas.
- Propuesta válida: «Chocolate de cacao ecuatoriano para regalar». Acción posible: «Explorar regalos» hacia una selección relevante. No exigir que repitan literalmente el ejemplo.
- Boceto: valorar relaciones, orden y adaptación, no habilidad artística. Si proponen testimonios o sellos, deben marcarlos como pendientes de evidencia; no inventarlos.
- Accesibilidad: ejemplos verificables: recorrido por Tab/Shift+Tab con foco visible; contraste medido contra su fondo; texto alternativo informativo; revisar que el contenido no exija desplazamiento horizontal en la vista propuesta.
- Cálculo: 50 / 2.000 × 100 = 2,5 %. La acción es una solicitud de información; no una venta. Las sesiones no equivalen a personas únicas. Comparar períodos sin controlar otras variaciones no demuestra causalidad del diseño.
- Rúbrica: máximo 20 puntos. No confundir esta escala propuesta con la distribución oficial de asignatura.
- Preguntas de Educaplay: clave y explicaciones completas en Actividad_Educaplay.txt.

## Límites
Minga Cacao y sus cifras son ficticios. No hay cuentas, campañas ni comercio real asociados. El material explica criterios; no certifica accesibilidad ni resultados comerciales.
''',encoding='utf-8')

def make_quiz():
    lines=['ACTIVIDAD PREPARADA PARA EDUCAPLAY — NO PUBLICADA','Título: Las páginas de inicio: orienta y decide','Tipo propuesto: cuestionario de opción múltiple (Quiz)','Materia: Negocios Electrónicos · U1 · 1.1','Objetivo: reconocer elementos, aplicar criterios de jerarquía y accesibilidad e interpretar una medida de resultado.','Instrucciones: lee cada situación y elige una sola respuesta. Revisa la retroalimentación antes de continuar.','Configuración sugerida: 10 preguntas, 1 punto por acierto, sin descuento por error, sin límite de tiempo. Referencia de logro: 8/10. Son decisiones didácticas propuestas.','Estado: banco de preguntas completo; no se accedió a una cuenta de Educaplay ni se publicó una actividad. Enlace real: pendiente de publicación.','']
    for i,q in enumerate(QUIZ,1):
        lines.append(f'{i}. {q["q"]}')
        for j,a in enumerate(q['options']):lines.append(f'   {"ABCD"[j]}) {a}')
        lines.append('Respuesta correcta: '+'ABCD'[q['answer']]);lines.append('Retroalimentación: '+q['feedback']);lines.append('')
    lines += ['PASOS PARA PUBLICAR CON TU CUENTA AUTORIZADA',
      '1. Inicia sesión en Educaplay desde su sitio oficial: https://www.educaplay.com/',
      '2. Elige crear una actividad y selecciona el formato de cuestionario de opciones múltiples disponible. La interfaz puede variar según cuenta y versión.',
      '3. Copia el título, objetivo e instrucciones y carga manualmente las 10 preguntas, sus cuatro opciones y la respuesta correcta.',
      '4. Añade la retroalimentación si el formato lo permite. Si no, conserva esta guía para la revisión docente. No confíes en una importación automática: el CSV adjunto es un banco auxiliar, no un formato de importación verificado de Educaplay.',
      '5. Configura español, nivel acorde a la asignatura y las opciones de evaluación que tu plan permita. Evita un cronómetro como barrera innecesaria.',
      '6. Revisa cada respuesta correcta en la vista de prueba y comprueba un intento completo antes de publicar.',
      '7. Publica según las opciones de privacidad de tu cuenta y copia el enlace real de compartir.',
      '8. Abre ese enlace con los permisos de un estudiante. Pégalo aquí únicamente después de comprobarlo. No uses una URL de ejemplo.',
      '', 'VERSIÓN LOCAL: Actividad_Interactiva.html permite practicar estas mismas preguntas sin cuenta ni conexión. No es una actividad alojada en Educaplay.',
      'Fuente: elaboración didáctica original; sílabo UPS de Negocios Electrónicos, tema 1.1; WCAG 2.2 para el criterio de contraste.']
    (TOPIC/'Actividad_Educaplay.txt').write_text('\n'.join(lines),encoding='utf-8')
    with (TOPIC/'Recursos'/'Banco_preguntas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['numero','pregunta','A','B','C','D','correcta','retroalimentacion'])
        for i,q in enumerate(QUIZ,1):writer.writerow([i,q['q'],*q['options'],'ABCD'[q['answer']],q['feedback']])
    document='''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Práctica · Las páginas de inicio</title>
<style>*{box-sizing:border-box}body{margin:0;background:#f5f2eb;color:#18394a;font:18px/1.6 system-ui,sans-serif}main{max-width:860px;margin:3rem auto;padding:1.8rem;background:white;border-radius:20px;border-top:8px solid #087e8b}h1{line-height:1.15}small{color:#536975}button{font:inherit;cursor:pointer}#opciones button{display:block;width:100%;text-align:left;margin:12px 0;padding:16px;border:2px solid #bccdcc;border-radius:12px;background:#fff;color:#18394a}button:focus-visible{outline:4px solid #b9432e;outline-offset:3px}#opciones button:disabled{cursor:default;opacity:1}#siguiente,#reiniciar{padding:12px 24px;background:#087e8b;color:white;border:0;border-radius:10px;margin-top:16px}#respuesta{padding:16px;border-radius:10px}progress{width:100%;accent-color:#087e8b}a{color:#087e8b}@media(max-width:600px){main{margin:0;padding:1.2rem}}</style>
<main><small>NEGOCIOS ELECTRÓNICOS · TEMA 1.1</small><h1>Las páginas de inicio</h1><p>Elige una respuesta por pregunta. Revisa la explicación y continúa. Sin límite de tiempo.</p><p><small>Práctica local independiente. No publicada en Educaplay. No recopila datos ni usa servicios externos.</small></p><progress id="progreso" value="0" max="10" aria-label="Preguntas completadas"></progress><p id="contador"></p><h2 id="pregunta" tabindex="-1"></h2><div id="opciones"></div><div id="respuesta" role="status" aria-live="polite"></div><button id="siguiente" hidden>Siguiente pregunta</button><button id="reiniciar" hidden>Volver a practicar</button><noscript>Activa JavaScript para practicar. El banco completo también está disponible en Actividad_Educaplay.txt.</noscript></main>
<script>const preguntas=__DATA__;let indice=0,puntos=0,bloqueado=false;const $=id=>document.getElementById(id);
function mostrar(){bloqueado=false;$('contador').textContent=`Pregunta ${indice+1} de ${preguntas.length}`;$('pregunta').textContent=preguntas[indice].q;$('opciones').replaceChildren();$('respuesta').textContent='';$('respuesta').style.background='transparent';$('siguiente').hidden=true;preguntas[indice].options.forEach((texto,i)=>{const b=document.createElement('button');b.textContent=texto;b.onclick=()=>responder(i);$('opciones').append(b)});}
function responder(i){if(bloqueado)return;bloqueado=true;const q=preguntas[indice],bien=i===q.answer;if(bien)puntos++;[...$('opciones').children].forEach((b,j)=>{b.disabled=true;if(j===q.answer){b.style.borderColor='#087e8b';b.textContent+=' ✓ Respuesta correcta'}else if(j===i)b.textContent+=' — Tu elección'});$('respuesta').style.background=bien?'#e2f0ee':'#fff0e8';$('respuesta').textContent=(bien?'Correcto. ':'Revisa la respuesta. ')+q.feedback;$('progreso').value=indice+1;$('siguiente').hidden=false;$('siguiente').textContent=indice===preguntas.length-1?'Ver resultado':'Siguiente pregunta';}
$('siguiente').onclick=()=>{indice++;if(indice<preguntas.length){mostrar();$('pregunta').focus()}else{$('contador').textContent='Práctica completada';$('pregunta').textContent=`Resultado: ${puntos} de ${preguntas.length}`;$('opciones').replaceChildren();$('respuesta').textContent=puntos>=8?'Has alcanzado la referencia de práctica. Revisa también cómo aplicarías los criterios.':'Vuelve a la presentación y revisa las explicaciones antes de repetir.';$('siguiente').hidden=true;$('reiniciar').hidden=false;$('pregunta').focus()}};
$('reiniciar').onclick=()=>{indice=0;puntos=0;$('progreso').value=0;$('reiniciar').hidden=true;mostrar();$('pregunta').focus()};mostrar();</script></html>'''
    (TOPIC/'Actividad_Interactiva.html').write_text(document.replace('__DATA__',json.dumps(QUIZ,ensure_ascii=False)),encoding='utf-8')

def make_plan():
    rows=[]
    for m in SYLLABI:
        for ui,(unit,topics) in enumerate(m['unidades'],1):
            for num,title in topics:
                pilot=m['codigo']=='E-CA-LND-202' and num=='1.1'
                rows.append([m['materia'],f'{ui}. {unit}',num,title,'PPTX · MP4 · MP3 · DOCX · Educaplay','Piloto propuesto' if pilot else 'Asignación por confirmar'])
    with (ROOT/'Matriz_temas.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['Materia','Unidad','Subtema','Contenido literal (sin puntuación final)','Materiales requeridos','Alcance']);w.writerows(rows)
    md=['# Análisis de los sílabos y matriz de producción','',
    'Fuente: los dos PDF aportados por el usuario, Universidad Politécnica Salesiana, período 2026–2027; docente Edgar Olmedo Andrade Espín. Se conserva la secuencia de contenidos de la página 1 de cada sílabo; solo se normalizan espacios y puntuación final. Los encabezados de unidad se mantienen. Cada subtema numerado es una unidad de producción propuesta, no una asignación confirmada por el docente.','',
    '## Materias','', '| Materia | Código | Nivel | Horas | Unidades | Subtemas |', '|---|---|---:|---:|---:|---:|',
    '| Negocios Electrónicos | E-CA-LND-202 | 4 | 120 | 4 | 12 |', '| Procesos para Negocios Digitales | E-CA-LND-214 | 8 | 160 | 4 | 15 |','',
    '## Tabla completa de materiales','', '| Materia | Unidad | Nº | Tema / subtema literal | Materiales | Alcance |','|---|---|---|---|---|---|']
    for row in rows:md.append('| '+' | '.join(row)+' |')
    md += ['', 'Son 27 subtemas y 135 entregables si se seleccionan todos. Una actividad preparada para Educaplay no equivale a una actividad publicada.', '',
    '## Resultados de aprendizaje de las unidades','',
    '**Negocios Electrónicos:** U1: Conoce la estructura de un negocio electrónico. U2: Aplica los principios y las técnicas de gestión de calidad y de innovación tecnológica en las organizaciones. U3: Aplica sistemas basados en Internet (ecommerce, e-learning, etc). Desarrollar soluciones de negocio mediante la implantación y la integración de hardware y software. U4: Diseña estrategias para desarrollar negocios globales.', '',
    '**Procesos para Negocios Digitales:** U1: Reconoce los conceptos relacionados con BPM. U2: Identifica claramente la estructura de un BPM. U3: Maneja las técnicas de modelado de procesos. U4: Implementa proyectos con base a BPM.', '',
    '## Actividades oficiales y elementos que no son nuevos temas', '',
    'Negocios Electrónicos: el PDF tiene dos páginas; el apartado 3 aparece sin un desarrollo detallado de actividades en el archivo suministrado. No se han inventado tareas oficiales ni ponderaciones.', '',
    'Procesos para Negocios Digitales: pp. 2–6 incluyen trabajo/taller colaborativo, Tarea 1, Proyecto entregable 1, Tarea 2, Entregable 2 proyecto, dos talleres colaborativos de la U3, Entregable 3 proyecto y Entregable 4, además de los dos exámenes parciales. La tabla asigna 15 puntos a U1, 15 a U2, 20 a U3, 10 a U4 y 20 a cada examen: total 100. Las tareas, talleres y exámenes son actividades y evidencias, no subtemas adicionales. Los indicadores destacan vocabulario y ciclo BPM (U1), componentes e interpretación de esquemas (U2), notación BPMN (U3) y propuestas, herramientas y KPI (U4).', '',
    'En la planificación aparece una variante truncada de 3.1 en Procesos para Negocios Digitales; se usa como denominación canónica la estructura de la página 1: «Importancia del modelado en la gestión de procesos». Las bibliografías y el perfil docente se conservan en los PDF originales, incluidos en Fuentes/.', '',
    '## Fase 2: selección pendiente', '',
    'Los documentos no identifican al estudiante ni le asignan contenidos particulares. Todos los subtemas de la tabla son candidatos de selección; corresponde confirmar la asignación con el profesor. Para avanzar se produjo como propuesta el subtema 1.1 de Negocios Electrónicos: Las páginas de inicio. Esta selección no se presenta como obligación del docente. Los otros 26 subtemas no se han desarrollado.', '',
    'Propuestas de alcance, si el profesor permite elegir: un tema de entrada (NE 1.1); un tema introductorio de la segunda materia (PND 1.1); o una unidad completa, previa confirmación. No se atribuye ninguna de estas propuestas al sílabo.', '',
    '## Procedencia', '', *[f'- {a} {b}' for a,b in SOURCES]]
    (ROOT/'Analisis_silabos.md').write_text('\n'.join(md),encoding='utf-8')
    d=setup_doc('Matriz de temas de los dos sílabos');d.add_heading('Plan de producción educativa',0)
    d.add_paragraph('27 subtemas · 2 materias · 5 entregables por subtema seleccionado. Asignación individual pendiente. Piloto propuesto: NE 1.1. Las páginas de inicio.')
    for m in SYLLABI:
        d.add_heading(m['materia'],1);d.add_paragraph(f"{m['codigo']} · Nivel {m['nivel']} · {m['horas']} horas · Período 2026–2027")
        for ui,(unit,topics) in enumerate(m['unidades'],1):
            d.add_heading(f'Unidad {ui}. {unit}',2)
            table(d,['Nº','Tema','Entregables necesarios'],[[num,title,'PPTX, MP4, MP3, DOCX y Educaplay'] for num,title in topics])
    d.add_paragraph('Fuente: página 1 de cada sílabo adjunto. Ver Analisis_silabos.md para resultados de aprendizaje, actividades oficiales y alcance. No se omiten temas de la matriz; solo el piloto está desarrollado.')
    d.save(ROOT/'Matriz_temas.docx')
    (ROOT/'Fuentes').mkdir(exist_ok=True)
    source_dir=Path(__import__('os').environ.get('EDU_SOURCES_DIR','/workspace/attachments'))
    if not source_dir.exists():source_dir=ROOT/'Fuentes'
    for p in source_dir.rglob('*.pdf'):
        dest=ROOT/'Fuentes'/p.name
        if p.resolve()!=dest.resolve():shutil.copy2(p,dest)
    (ROOT/'Procesos_para_Negocios_Digitales').mkdir(exist_ok=True)
    (ROOT/'Procesos_para_Negocios_Digitales'/'SELECCION_PENDIENTE.md').write_text('La matriz contiene los 15 subtemas exactos de esta materia. No se han generado materiales para esta asignatura: falta confirmar qué subtemas corresponden al estudiante. Ver ../Analisis_silabos.md.\n',encoding='utf-8')

def save_scripts():
    (TOPIC/'Recursos'/'Guion_video.md').write_text('# Guion del video · tema 1.1\n\n'+'\n\n'.join(f'## Escena {i+1}: {s["title"].replace(chr(10)," ")}\n\n{s["narration"]}' for i,s in enumerate(SLIDES)),encoding='utf-8')
    (TOPIC/'Recursos'/'Guion_audio.md').write_text('# Audio educativo independiente · Las páginas de inicio\n\n'+'\n\n'.join(AUDIO_PARAGRAPHS),encoding='utf-8')

if __name__=='__main__':
    ROOT.mkdir(parents=True,exist_ok=True)
    make_slides();make_task();make_quiz();make_plan();save_scripts()
